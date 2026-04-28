from __future__ import annotations

import importlib.util
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SKILL_DIR = ROOT / "skills" / "thesis-word-format"
SCRIPT_PATH = SKILL_DIR / "scripts" / "normalize_citation_markers.py"


def load_normalizer():
    spec = importlib.util.spec_from_file_location("normalize_citation_markers", SCRIPT_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def test_skill_bundle_has_required_guidance_files():
    assert (SKILL_DIR / "SKILL.md").exists()
    assert (SKILL_DIR / "references" / "degree-format-summary.md").exists()
    assert (SKILL_DIR / "references" / "citation-position-rules.md").exists()
    assert (SKILL_DIR / "references" / "docx-format-execution-checklist.md").exists()
    assert SCRIPT_PATH.exists()

    skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
    assert "博士学位论文" in skill_text
    assert "文献[n]" in skill_text
    assert "上标" in skill_text
    assert "docx-format-execution-checklist.md" in skill_text
    assert "CSL" in skill_text
    assert "metadata" in skill_text
    assert "postprocess" in skill_text
    assert "语义" in skill_text


def test_docx_format_checklist_covers_execution_and_validation():
    checklist = (SKILL_DIR / "references" / "docx-format-execution-checklist.md").read_text(encoding="utf-8")

    for required in [
        "Word style mapping",
        "A4",
        "160 mm x 247 mm",
        "Heading 1",
        "页眉",
        "页脚",
        "参考文献",
        "LibreOffice",
        "Poppler",
        "risk boundary",
    ]:
        assert required in checklist


def test_citation_marker_context_classification():
    normalizer = load_normalizer()

    assert normalizer.classify_marker("已有研究指出[12]。", 6, 10) == "terminal"
    assert normalizer.classify_marker("文献[12]指出该方法有效。", 2, 6) == "terminal"
    assert normalizer.classify_marker("由文献[8,10-14]可知，该问题仍存在。", 3, 12) == "terminal"


def test_english_author_led_citation_is_reported_for_chinese_rewrite():
    normalizer = load_normalizer()

    issues = normalizer.scan_paragraph_issues("Sun等[1]系统梳理了算力网络。")
    assert issues
    assert issues[0]["kind"] == "english_author_led_citation"
    assert "文献[1]" in issues[0]["suggestion"]

    assert normalizer.scan_paragraph_issues("文献[1]系统梳理了算力网络。") == []


def test_english_author_led_citations_are_rewritten_to_wenxian():
    normalizer = load_normalizer()

    rewritten, rewrites = normalizer.rewrite_author_led_citations(
        "Chen等[51]在理论层面给出了结果。随后，Geng等[52]研究传统路由系统。Büsing等[53]处理鲁棒约束。"
    )
    assert rewritten == "文献[51]在理论层面给出了结果。随后，文献[52]研究传统路由系统。文献[53]处理鲁棒约束。"
    assert [item["marker"] for item in rewrites] == ["[51]", "[52]", "[53]"]

    unchanged, rewrites = normalizer.rewrite_author_led_citations("该方法继承了Chen等[51]的模型。")
    assert unchanged == "该方法继承了Chen等[51]的模型。"
    assert rewrites == []


def test_docx_normalization_sets_terminal_markers_as_superscript(tmp_path):
    normalizer = load_normalizer()
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "output.docx"

    doc = Document()
    doc.add_paragraph("已有研究指出[1]。")
    doc.add_paragraph("文献[2]指出该方法有效。")
    doc.save(input_path)

    report = normalizer.normalize_docx(input_path, output_path)

    result = Document(output_path)
    terminal_marker_runs = [
        run for run in result.paragraphs[0].runs if "[1]" in run.text
    ]
    narrative_marker_runs = [
        run for run in result.paragraphs[1].runs if "[2]" in run.text
    ]
    assert terminal_marker_runs and terminal_marker_runs[0].font.superscript is True
    assert narrative_marker_runs and narrative_marker_runs[0].font.superscript is True
    assert report["styled_markers"] == 2


def test_docx_normalization_preserves_existing_run_formatting(tmp_path):
    normalizer = load_normalizer()
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "output.docx"

    doc = Document()
    paragraph = doc.add_paragraph()
    prefix = paragraph.add_run("关键结论")
    prefix.bold = True
    paragraph.add_run("已有研究指出[1]。")
    doc.save(input_path)

    normalizer.normalize_docx(input_path, output_path)

    result = Document(output_path)
    assert result.paragraphs[0].runs[0].text == "关键结论"
    assert result.paragraphs[0].runs[0].bold is True


def test_docx_normalization_handles_table_cell_markers(tmp_path):
    normalizer = load_normalizer()
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "output.docx"

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格结论[3]。"
    doc.save(input_path)

    report = normalizer.normalize_docx(input_path, output_path)

    result = Document(output_path)
    runs = result.tables[0].cell(0, 0).paragraphs[0].runs
    marker_runs = [run for run in runs if "[3]" in run.text]
    assert marker_runs and marker_runs[0].font.superscript is True
    assert report["styled_markers"] == 1


def test_docx_normalization_rewrites_english_author_led_citation(tmp_path):
    normalizer = load_normalizer()
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "output.docx"

    doc = Document()
    doc.add_paragraph("前文已有说明。Chen等[51]在理论层面给出了结果。")
    doc.save(input_path)

    report = normalizer.normalize_docx(input_path, output_path)

    result = Document(output_path)
    paragraph = result.paragraphs[0]
    assert paragraph.text == "前文已有说明。文献[51]在理论层面给出了结果。"
    marker_runs = [run for run in paragraph.runs if "[51]" in run.text]
    assert marker_runs and marker_runs[0].font.superscript is True
    assert report["rewritten_author_led_citations"] == 1


def test_author_led_rewrite_preserves_existing_run_formatting(tmp_path):
    normalizer = load_normalizer()
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "output.docx"

    doc = Document()
    paragraph = doc.add_paragraph()
    prefix = paragraph.add_run("前文已有说明。")
    prefix.bold = True
    paragraph.add_run("Chen等[51]在理论层面给出了结果。")
    doc.save(input_path)

    normalizer.normalize_docx(input_path, output_path)

    result = Document(output_path)
    runs = result.paragraphs[0].runs
    assert runs[0].text == "前文已有说明。"
    assert runs[0].bold is True
    assert runs[1].text == "文献"
    assert runs[1].bold in (False, None)
    marker_runs = [run for run in runs if "[51]" in run.text]
    assert marker_runs and marker_runs[0].font.superscript is True
