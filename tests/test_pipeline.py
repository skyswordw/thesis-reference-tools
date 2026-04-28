import pytest
from docx import Document

from thesis_refs.pipeline import (
    audit_zotero_docx,
    build_reference_model,
    bibtex_author_value,
    compress_citation_runs,
    expand_citation_cluster,
    generate_zotero_migration_checklist,
    normalize_space,
    optimize_docx_media,
    postprocess_zotero_bibliography_docx,
    rewrite_main_docx,
    make_dedupe_key,
    prepare_zotero_operator_docx,
    replace_nth_citation_cluster,
    write_zotero_group_handoff,
    write_zotero_migration_checklist,
    load_pipeline_config,
)


@pytest.fixture()
def demo_docx(tmp_path):
    doc = Document()
    doc.add_heading("第一章", level=1)
    doc.add_heading("1.1 Demo section", level=2)
    doc.add_paragraph("第一章引用本节文献[1][2]。")
    doc.add_paragraph("[1] Demo Org. Alpha sensing route[J]. Demo Journal, 2024.")
    doc.add_paragraph(
        "[2] ITU-T. Q.4144: Signalling requirements for cross-operator service "
        "orchestration in computing power networks[EB/OL]. 2025-04-13."
    )

    doc.add_heading("第二章", level=1)
    doc.add_paragraph("第二章引用[1][2]，并复用[3]。")
    doc.add_paragraph("组合引用[2][3]需要保留顺序。")
    doc.add_paragraph("[1] Demo Lab. Beta routing[J]. Demo Journal, 2025.")
    doc.add_paragraph("[2] Demo Alliance. Gamma compute[R/OL]. 2024.")
    doc.add_paragraph(
        "[3] ITU-T. Q.4144: Signalling requirements for cross-operator service "
        "orchestration in computing power networks[EB/OL]. 2025-04-13."
    )

    doc.add_heading("参考文献", level=1)
    doc.add_heading("致谢", level=1)

    path = tmp_path / "demo-single-docx.docx"
    doc.save(path)
    return path


def test_extracts_references_from_single_demo_document(demo_docx):
    model = build_reference_model(demo_docx)

    chapter_one = [ref for ref in model.references if ref.source_chapter == "第一章"]
    chapter_two = [ref for ref in model.references if ref.source_chapter == "第二章"]

    assert len(chapter_one) == 2
    assert {ref.source_section for ref in chapter_one} == {"1.1 Demo section"}
    assert len(chapter_two) == 3
    assert len(model.citation_occurrences) == 7


def test_pipeline_config_accepts_only_single_doc_settings(tmp_path):
    config = tmp_path / "config.toml"
    config.write_text(
        'input_docx = "examples/demo/raw/thesis.docx"\noutput_root = "examples/demo/output"\n',
        encoding="utf-8",
    )

    assert load_pipeline_config(config) == {
        "input_docx": "examples/demo/raw/thesis.docx",
        "output_root": "examples/demo/output",
    }

    bad_config = tmp_path / "bad.toml"
    bad_config.write_text('input_docx = "x.docx"\n' + "reference_" + "sources = []\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported config keys"):
        load_pipeline_config(bad_config)


def test_known_duplicate_references_share_one_key():
    q4144_variants = [
        "ITU-T. Q.4144: Signalling requirements for cross-operator service orchestration in computing power networks[EB/OL]. 2025-04-13.",
        "[6] ITU-T. Q.4144: Signalling requirements for cross-operator service orchestration in computing power networks[EB/OL]. 2025-04-13.",
    ]

    assert make_dedupe_key(q4144_variants[0]) == make_dedupe_key(q4144_variants[1])


def test_global_numbering_merges_demo_duplicate(demo_docx):
    model = build_reference_model(demo_docx)
    q4144_refs = [
        ref for ref in model.references
        if "Q.4144" in ref.raw_text and "Signalling requirements" in ref.raw_text
    ]

    assert len(q4144_refs) == 2
    assert len({model.global_id_by_ref_id[ref.ref_id] for ref in q4144_refs}) == 1


def test_compresses_adjacent_citation_runs():
    assert compress_citation_runs("现有研究[1][2][3]指出") == "现有研究[1-3]指出"
    assert compress_citation_runs("现有研究[1][3][4]指出") == "现有研究[1][3-4]指出"


def test_bibtex_author_value_is_zotero_friendly():
    assert bibtex_author_value("Sun Y, Lei B, Liu J, et al") == "Sun, Y and Lei, B and Liu, J and others"
    assert bibtex_author_value("柴若楠, 郜帅, 兰江雨等") == "柴若楠 and 郜帅 and 兰江雨 and others"
    assert bibtex_author_value("中共中央, 国务院") == "{{中共中央}} and {{国务院}}"


def test_rebuilt_docx_has_ordered_final_references_and_no_chapter_one_local_blocks(tmp_path, demo_docx):
    model = build_reference_model(demo_docx)
    output = rewrite_main_docx(model, tmp_path / "unified.docx", demo_docx)
    doc = Document(output)

    ref_heading = next(
        idx for idx, paragraph in enumerate(doc.paragraphs)
        if paragraph.style.name == "Heading 1" and paragraph.text.strip() == "参考文献"
    )
    final_refs = []
    for paragraph in doc.paragraphs[ref_heading + 1:]:
        if paragraph.style.name == "Heading 1":
            break
        text = paragraph.text.strip()
        if text:
            final_refs.append(text)

    assert len(final_refs) == len(model.unique_references)
    assert final_refs[0].startswith("[1] ")
    assert final_refs[-1].startswith(f"[{len(model.unique_references)}] ")

    first_chapter_end = next(
        idx for idx, paragraph in enumerate(doc.paragraphs)
        if paragraph.style.name == "Heading 1" and paragraph.text.strip().startswith("第二章")
    )
    first_chapter_ref_blocks = [
        paragraph.text.strip() for paragraph in doc.paragraphs[:first_chapter_end]
        if paragraph.text.strip().startswith("[")
    ]
    assert first_chapter_ref_blocks == []


def test_zotero_migration_checklist_groups_static_citation_positions(tmp_path, demo_docx):
    model = build_reference_model(demo_docx)
    unified_docx = rewrite_main_docx(model, tmp_path / "unified.docx", demo_docx)
    unified_doc = Document(unified_docx)
    checklist = generate_zotero_migration_checklist(model, unified_docx)

    assert checklist["summary"]["unique_references"] == len(model.unique_references)
    assert checklist["summary"]["covered_citation_occurrences"] == len(model.citation_occurrences)
    assert checklist["summary"]["paragraph_coverage_ok"] is True
    assert checklist["coverage_mismatches"] == []
    assert checklist["summary"]["citation_positions"] >= len(
        {occ.paragraph_index for occ in model.citation_occurrences}
    )

    first = checklist["positions"][0]
    assert first["marker_text"] == "[1-2]"
    assert first["global_ids"] == [1, 2]
    assert first["zotero_keys"] == ["ref001", "ref002"]
    assert normalize_space(unified_doc.paragraphs[first["paragraph_index"]].text) == first["paragraph_text"]

    assert all(
        normalize_space(unified_doc.paragraphs[position["paragraph_index"]].text) == position["paragraph_text"]
        and position["marker_text"] in position["paragraph_text"]
        for position in checklist["positions"]
    )

    assert any(
        position["marker_text"] == "[3-4]"
        and position["global_ids"] == [3, 4]
        and position["zotero_keys"] == ["ref003", "ref004"]
        for position in checklist["positions"]
    )
    assert any(
        position["marker_text"] == "[4][2]" and position["global_ids"] == [4, 2]
        for position in checklist["positions"]
    )

    json_path, md_path = write_zotero_migration_checklist(model, tmp_path, unified_docx)
    assert json_path.exists()
    assert md_path.exists()
    assert "Zotero 动态域迁移清单" in md_path.read_text(encoding="utf-8")


def test_expand_citation_cluster_handles_ranges_and_adjacent_markers():
    assert expand_citation_cluster("[27-29]") == [27, 28, 29]
    assert expand_citation_cluster("[3][6-7]") == [3, 6, 7]


def test_audit_zotero_docx_counts_real_fields_without_flagging_field_results(tmp_path):
    import zipfile

    sample = tmp_path / "sample.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:fldSimple w:instr=" ADDIN ZOTERO_ITEM CSL_CITATION {} "><w:r><w:t>[9]</w:t></w:r></w:fldSimple></w:p>
    <w:p>
      <w:r><w:fldChar w:fldCharType="begin"/></w:r>
      <w:r><w:instrText> ADDIN ZOTERO_ITEM CSL_CITATION {} </w:instrText></w:r>
      <w:r><w:fldChar w:fldCharType="separate"/></w:r>
      <w:r><w:t>[8]</w:t></w:r>
      <w:r><w:fldChar w:fldCharType="end"/></w:r>
    </w:p>
    <w:p><w:r><w:instrText> ADDIN ZOTERO_BIBL CSL_BIBLIOGRAPHY </w:instrText></w:r></w:p>
    <w:p><w:r><w:t>仍未迁移的静态引用[1]</w:t></w:r></w:p>
    <w:p><w:r><w:t>另一个段落里的静态引用[2]</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)

    audit = audit_zotero_docx(sample, tmp_path)

    assert audit["zotero_item_fields"] == 2
    assert audit["csl_citation_fields"] == 2
    assert audit["csl_bibliography_fields"] == 1
    assert audit["remaining_static_citation_markers"] == 2
    assert (tmp_path / "zotero_field_audit.md").exists()


def test_audit_zotero_docx_counts_bibliography_field_before_result_end(tmp_path):
    import zipfile

    sample = tmp_path / "bibliography.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p>
      <w:r><w:fldChar w:fldCharType="begin"/></w:r>
      <w:r><w:instrText> ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY </w:instrText></w:r>
      <w:r><w:fldChar w:fldCharType="separate"/></w:r>
      <w:r><w:t>[1]\tAuthor A. Title[J]. Journal, 2024.</w:t></w:r>
    </w:p>
  </w:body>
</w:document>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)

    audit = audit_zotero_docx(sample, tmp_path)

    assert audit["csl_bibliography_fields"] == 1
    assert audit["complex_field_begin_count"] == 1
    assert audit["complex_field_end_count"] == 0
    assert audit["complex_field_unclosed_count"] == 1
    assert audit["bibliography_field_unclosed_count"] == 1


def test_audit_zotero_docx_keeps_adjacent_paragraph_markers_separate(tmp_path):
    import zipfile

    sample = tmp_path / "adjacent.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>[1]</w:t></w:r></w:p>
    <w:p><w:r><w:t>[2]</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)

    audit = audit_zotero_docx(sample, tmp_path)

    assert audit["remaining_static_citation_markers"] == 2
    assert audit["remaining_static_citation_samples"] == ["[1]", "[2]"]


def test_zotero_group_handoff_mentions_group_library_and_csl(tmp_path, demo_docx):
    model = build_reference_model(demo_docx)

    output = write_zotero_group_handoff(model, tmp_path)
    content = output.read_text(encoding="utf-8")

    assert "Group Library" in content
    assert "GB/T 7714-2015" in content
    assert "refs.bib" in content
    assert str(len(model.unique_references)) in content


def test_prepare_zotero_operator_docx_replaces_static_citations_with_unique_placeholders(tmp_path, demo_docx):
    model = build_reference_model(demo_docx)
    unified_docx = rewrite_main_docx(model, tmp_path / "unified.docx", demo_docx)

    operator_docx, checklist_json, checklist_md = prepare_zotero_operator_docx(
        model,
        source_docx_path=unified_docx,
        output_docx_path=tmp_path / "operator.docx",
        output_dir=tmp_path,
    )

    full_text = "\n".join(paragraph.text for paragraph in Document(operator_docx).paragraphs)

    assert "[[ZOTERO_P001]]" in full_text
    assert "[[ZOTERO_P004]]" in full_text
    assert "[4][2]" not in full_text
    assert checklist_json.exists()
    assert checklist_md.exists()

    audit = audit_zotero_docx(operator_docx, tmp_path, checklist_json)
    checklist = generate_zotero_migration_checklist(model, unified_docx)
    assert audit["operator_placeholders"] == checklist["summary"]["citation_positions"]
    assert audit["operator_placeholder_sequence_matches"] is True


def test_replace_nth_citation_cluster_handles_split_runs_without_rewriting_paragraph(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("前文")
    paragraph.add_run("[")
    paragraph.add_run("12")
    paragraph.add_run("-")
    paragraph.add_run("13")
    paragraph.add_run("]")
    paragraph.add_run("中间")
    paragraph.add_run("[2]")
    paragraph.add_run("后文")

    replaced = replace_nth_citation_cluster(paragraph, 1, "[@ref012; @ref013]", "[12-13]")
    replace_nth_citation_cluster(paragraph, 1, "[[ZOTERO_P002]]", "[2]")

    assert replaced == "[12-13]"
    assert paragraph.text == "前文[@ref012; @ref013]中间[[ZOTERO_P002]]后文"
    assert paragraph.runs[0].text == "前文"
    assert paragraph.runs[-1].text == "后文"


def test_audit_zotero_docx_extracts_citation_keys_from_real_field_json(tmp_path):
    import zipfile

    checklist = tmp_path / "checklist.json"
    checklist.write_text(
        """{
  "positions": [
    {"position_id": "P001", "zotero_keys": ["ref001", "ref002"]}
  ]
}
""",
        encoding="utf-8",
    )
    sample = tmp_path / "zotero-json.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p>
      <w:r><w:fldChar w:fldCharType="begin"/></w:r>
      <w:r><w:instrText> ADDIN ZOTERO_ITEM CSL_CITATION {"citationItems":[{"itemData":{"citation-key":"ref001"}},{"itemData":{"note":"citation-key: ref002"}}]} </w:instrText></w:r>
      <w:r><w:fldChar w:fldCharType="separate"/></w:r>
      <w:r><w:t>[1,2]</w:t></w:r>
      <w:r><w:fldChar w:fldCharType="end"/></w:r>
    </w:p>
  </w:body>
</w:document>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)

    audit = audit_zotero_docx(sample, tmp_path, checklist)

    assert audit["csl_citation_fields"] == 1
    assert audit["extracted_citation_items"] == 2
    assert audit["zotero_citation_keys"] == ["ref001", "ref002"]
    assert audit["zotero_citation_key_sequence_matches"] is True


def test_audit_zotero_docx_ignores_out_of_range_numeric_arrays_with_checklist(tmp_path):
    import zipfile

    checklist = tmp_path / "checklist.json"
    checklist.write_text(
        """{
  "positions": [
    {"position_id": "P001", "zotero_keys": ["ref001"]}
  ]
}
""",
        encoding="utf-8",
    )
    sample = tmp_path / "numeric-array.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>未迁移残留引用 [1]</w:t></w:r></w:p>
    <w:p><w:r><w:t>[128, 128, 64]</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)

    audit = audit_zotero_docx(sample, tmp_path, checklist)

    assert audit["remaining_static_citation_markers"] == 1
    assert audit["remaining_static_citation_samples"] == ["[1]"]


def test_audit_zotero_docx_flags_bibliography_format_residue_and_style(tmp_path):
    import zipfile

    sample = tmp_path / "bad-bibliography.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>参考文献</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Bibliography"/><w:ind w:firstLine="480"/><w:spacing w:line="240" w:lineRule="exact"/><w:jc w:val="both"/></w:pPr><w:r><w:t>[[ZOTERO_BIBLIOGRAPHY[90]HUANG X, LIU R R, WU J. STPer[C/OL]//GLOBECOM 2025. https://doi.org/10.1109/globecom59602.2025.11431840. DOI:10.1109/globecom59602.2025.11431840.</w:t></w:r></w:p>
    <w:p><w:r><w:t>]]</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>致谢</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)

    audit = audit_zotero_docx(sample, tmp_path)

    assert audit["bibliography_anchor_residue_count"] == 1
    assert audit["stray_closing_bracket_paragraphs"] == 1
    assert audit["bibliography_first_line_indent_count"] == 1
    assert audit["bibliography_tight_line_spacing_count"] == 1
    assert audit["bibliography_non_left_alignment_count"] == 1
    assert audit["bibliography_all_caps_author_samples"]
    assert audit["bibliography_online_article_conference_samples"]
    assert audit["bibliography_duplicate_doi_url_samples"]


def test_postprocess_zotero_bibliography_removes_residue_and_applies_hanging_indent(tmp_path):
    import zipfile
    from xml.etree import ElementTree

    sample = tmp_path / "postprocess-input.docx"
    output = tmp_path / "postprocess-output.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>参考文献</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Bibliography"/><w:ind w:firstLine="480"/><w:spacing w:line="240" w:lineRule="exact"/><w:jc w:val="both"/></w:pPr><w:r><w:t>[[ZOTERO_BIBLIOGRAPHY[90]Huang X, Liu R R, Wu J. STPer[C]//GLOBECOM 2025. DOI:10.1109/globecom59602.2025.11431840.</w:t></w:r></w:p>
    <w:p><w:r><w:t>]]</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>致谢</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    styles_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:styleId="Bibliography"><w:name w:val="Bibliography"/><w:pPr><w:ind w:firstLine="480"/><w:spacing w:line="240" w:lineRule="exact"/></w:pPr></w:style>
</w:styles>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/styles.xml", styles_xml)

    postprocess_zotero_bibliography_docx(sample, output)
    audit = audit_zotero_docx(output, tmp_path)
    with zipfile.ZipFile(output) as archive:
        processed_xml = archive.read("word/document.xml").decode("utf-8")
        processed_styles = archive.read("word/styles.xml").decode("utf-8")
    root = ElementTree.fromstring(processed_xml.encode("utf-8"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    ref_ppr = next(
        ppr
        for ppr in root.findall(".//w:pPr", ns)
        if (style := ppr.find("w:pStyle", ns)) is not None
        and style.attrib.get(f"{{{ns['w']}}}val") == "Bibliography"
    )
    assert ref_ppr is not None
    ref_ind = ref_ppr.find("w:ind", ns)
    ref_tab = ref_ppr.find("w:tabs/w:tab", ns)
    ref_spacing = ref_ppr.find("w:spacing", ns)
    ref_jc = ref_ppr.find("w:jc", ns)
    styles_root = ElementTree.fromstring(processed_styles.encode("utf-8"))
    style_ind = styles_root.find(".//w:style[@w:styleId='Bibliography']/w:pPr/w:ind", ns)

    assert audit["bibliography_anchor_residue_count"] == 0
    assert audit["stray_closing_bracket_paragraphs"] == 0
    assert audit["bibliography_first_line_indent_count"] == 0
    assert audit["bibliography_hanging_indent_count"] == 1
    assert audit["bibliography_tight_line_spacing_count"] == 0
    assert audit["bibliography_non_left_alignment_count"] == 0
    assert ref_ind is not None
    assert ref_ind.attrib[f"{{{ns['w']}}}left"] == "480"
    assert ref_ind.attrib[f"{{{ns['w']}}}hanging"] == "480"
    assert ref_tab is not None
    assert ref_tab.attrib[f"{{{ns['w']}}}val"] == "left"
    assert ref_tab.attrib[f"{{{ns['w']}}}pos"] == "480"
    assert ref_spacing is not None
    assert ref_spacing.attrib[f"{{{ns['w']}}}before"] == "0"
    assert ref_spacing.attrib[f"{{{ns['w']}}}after"] == "0"
    assert ref_spacing.attrib[f"{{{ns['w']}}}line"] == "300"
    assert ref_spacing.attrib[f"{{{ns['w']}}}lineRule"] == "auto"
    assert ref_jc is not None
    assert ref_jc.attrib[f"{{{ns['w']}}}val"] == "left"
    assert [child.tag.rsplit("}", 1)[-1] for child in ref_ppr] == [
        "pStyle",
        "tabs",
        "spacing",
        "ind",
        "jc",
        "rPr",
    ]
    assert style_ind is not None
    assert style_ind.attrib[f"{{{ns['w']}}}left"] == "480"
    assert style_ind.attrib[f"{{{ns['w']}}}hanging"] == "480"


def test_postprocess_zotero_bibliography_preserves_namespace_prefixes_used_by_mc_ignorable(tmp_path):
    import zipfile

    sample = tmp_path / "postprocess-prefix-input.docx"
    output = tmp_path / "postprocess-prefix-output.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" mc:Ignorable="w14">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>参考文献</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Bibliography"/></w:pPr><w:r><w:t>[1]\tAuthor A. Title[J]. Journal, 2024.</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    styles_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:styleId="Bibliography"><w:name w:val="Bibliography"/></w:style>
</w:styles>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/styles.xml", styles_xml)

    postprocess_zotero_bibliography_docx(sample, output)

    with zipfile.ZipFile(output) as archive:
        processed_xml = archive.read("word/document.xml").decode("utf-8")

    assert "xmlns:w14=" in processed_xml
    assert 'mc:Ignorable="w14"' in processed_xml
    assert "xmlns:ns" not in processed_xml.split(">", 1)[0]


def test_postprocess_zotero_bibliography_normalizes_latex_superscript_artifacts(tmp_path):
    import zipfile

    sample = tmp_path / "postprocess-latex-input.docx"
    output = tmp_path / "postprocess-latex-output.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>参考文献</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Bibliography"/></w:pPr><w:r><w:t>[58]\tZhang J. H$^{3}$DRA[J]. IEEE, 2026.</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    styles_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:styleId="Bibliography"><w:name w:val="Bibliography"/></w:style>
</w:styles>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/styles.xml", styles_xml)

    postprocess_zotero_bibliography_docx(sample, output)

    with zipfile.ZipFile(output) as archive:
        processed_xml = archive.read("word/document.xml").decode("utf-8")

    assert "H$^{3}$DRA" not in processed_xml
    assert "H³DRA" in processed_xml


def test_postprocess_zotero_bibliography_compacts_number_gap_and_hides_access_fields(tmp_path):
    import zipfile

    sample = tmp_path / "postprocess-access-input.docx"
    output = tmp_path / "postprocess-access-output.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>参考文献</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Bibliography"/></w:pPr><w:r><w:t>[11]</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>中国电信研究院. 云网一体信息基础设施系列白皮书[Z/OL]. (2023). https://www.chinatelecom.com.cn/ct/news/jczt/zxyw2024/kyzs/154975.html.</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Bibliography"/></w:pPr><w:r><w:t>[90]</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>Huang X, Liu R R, Wu J. STPer[C]//GLOBECOM 2025. DOI:10.1109/globecom59602.2025.11431840.</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    styles_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:styleId="Bibliography"><w:name w:val="Bibliography"/></w:style>
</w:styles>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/styles.xml", styles_xml)

    postprocess_zotero_bibliography_docx(sample, output)

    with zipfile.ZipFile(output) as archive:
        processed_xml = archive.read("word/document.xml").decode("utf-8")

    assert "[11] 中国电信研究院" in processed_xml
    assert "[90] Huang X" in processed_xml
    assert "<w:r><w:tab" not in processed_xml
    assert "https://www.chinatelecom.com.cn" not in processed_xml
    assert "DOI:10.1109" not in processed_xml


def test_audit_zotero_docx_flags_english_bibliography_using_chinese_et_al(tmp_path):
    import zipfile

    sample = tmp_path / "english-et-al-audit.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>参考文献</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Bibliography"/></w:pPr><w:r><w:t>[46] Yang D, Gong K, Ren J, 等. Intelligent Task Offloading[J]. IEEE Network, 2022, 36(2): 16-24.</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)

    audit = audit_zotero_docx(sample, tmp_path)

    assert audit["bibliography_english_chinese_et_al_count"] == 1
    assert "Yang D" in audit["bibliography_english_chinese_et_al_samples"][0]


def test_audit_zotero_docx_flags_report_and_standard_entries_falling_back_to_z(tmp_path):
    import json
    import zipfile

    sample = tmp_path / "semantic-type-audit.docx"
    checklist = tmp_path / "operator_migration_checklist.json"
    checklist.write_text(
        json.dumps(
            {
                "references": [
                    {"global_id": 8, "citekey": "ref008", "ref_type": "R/OL"},
                    {"global_id": 17, "citekey": "ref017", "ref_type": "S/OL"},
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>参考文献</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Bibliography"/></w:pPr><w:r><w:t>[8] 中国信息通信研究院. 中国综合算力评价白皮书（2023年）[Z/OL]. 2023.</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Bibliography"/></w:pPr><w:r><w:t>[17] ITU-T. Recommendation ITU-T Q.4144[S/OL]. 2025.</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)

    audit = audit_zotero_docx(sample, tmp_path, checklist)

    assert audit["bibliography_report_standard_z_type_count"] == 1
    assert "[8]" in audit["bibliography_report_standard_z_type_samples"][0]


def test_postprocess_zotero_bibliography_preserves_field_end_when_removing_closing_residue(tmp_path):
    import zipfile
    from xml.etree import ElementTree

    sample = tmp_path / "postprocess-field-end-input.docx"
    output = tmp_path / "postprocess-field-end-output.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>参考文献</w:t></w:r></w:p>
    <w:p>
      <w:pPr><w:pStyle w:val="Bibliography"/></w:pPr>
      <w:r><w:fldChar w:fldCharType="begin"/></w:r>
      <w:r><w:instrText> ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY </w:instrText></w:r>
      <w:r><w:fldChar w:fldCharType="separate"/></w:r>
      <w:r><w:t>[90]\tHuang X, Liu R R, Wu J. STPer[C]//GLOBECOM 2025. DOI:10.1109/globecom59602.2025.11431840.</w:t></w:r>
    </w:p>
    <w:p><w:r><w:fldChar w:fldCharType="end"/></w:r><w:r><w:t>]]</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>致谢</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    styles_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:styleId="Bibliography"><w:name w:val="Bibliography"/></w:style>
</w:styles>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/styles.xml", styles_xml)

    postprocess_zotero_bibliography_docx(sample, output)
    audit = audit_zotero_docx(output, tmp_path)
    with zipfile.ZipFile(output) as archive:
        processed_xml = archive.read("word/document.xml").decode("utf-8")
    root = ElementTree.fromstring(processed_xml.encode("utf-8"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    field_runs = [
        run
        for run in root.findall(".//w:r", ns)
        if run.find("w:fldChar", ns) is not None or run.find("w:instrText", ns) is not None
    ]

    assert audit["bibliography_anchor_residue_count"] == 0
    assert audit["stray_closing_bracket_paragraphs"] == 0
    assert audit["complex_field_begin_count"] == 1
    assert audit["complex_field_end_count"] == 1
    assert audit["complex_field_unclosed_count"] == 0
    assert audit["bibliography_field_unclosed_count"] == 0
    assert all(run[0].tag == f"{{{ns['w']}}}rPr" for run in field_runs if run.find("w:rPr", ns) is not None)


def test_postprocess_zotero_bibliography_repairs_missing_bibliography_field_end(tmp_path):
    import zipfile

    sample = tmp_path / "postprocess-missing-field-end-input.docx"
    output = tmp_path / "postprocess-missing-field-end-output.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>参考文献</w:t></w:r></w:p>
    <w:p>
      <w:pPr><w:pStyle w:val="Bibliography"/></w:pPr>
      <w:r><w:fldChar w:fldCharType="begin"/></w:r>
      <w:r><w:instrText> ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY </w:instrText></w:r>
      <w:r><w:fldChar w:fldCharType="separate"/></w:r>
      <w:r><w:t>[1]\tAuthor A. Title[J]. Journal, 2024.</w:t></w:r>
    </w:p>
    <w:p><w:pPr><w:pStyle w:val="Bibliography"/></w:pPr><w:r><w:t>[2]\tAuthor B. Title[J]. Journal, 2025.</w:t></w:r></w:p>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>致谢</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    styles_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:styleId="Bibliography"><w:name w:val="Bibliography"/></w:style>
</w:styles>
"""
    with zipfile.ZipFile(sample, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/styles.xml", styles_xml)

    postprocess_zotero_bibliography_docx(sample, output)
    audit = audit_zotero_docx(output, tmp_path)

    assert audit["csl_bibliography_fields"] == 1
    assert audit["complex_field_begin_count"] == 1
    assert audit["complex_field_end_count"] == 1
    assert audit["complex_field_unclosed_count"] == 0
    assert audit["bibliography_field_unclosed_count"] == 0


def test_optimize_docx_media_restores_source_media_without_losing_fields(tmp_path):
    import zipfile

    input_docx = tmp_path / "input.docx"
    media_source_docx = tmp_path / "media-source.docx"
    output_docx = tmp_path / "output.docx"
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p>
      <w:r><w:fldChar w:fldCharType="begin"/></w:r>
      <w:r><w:instrText> ADDIN ZOTERO_ITEM CSL_CITATION {} </w:instrText></w:r>
      <w:r><w:fldChar w:fldCharType="separate"/></w:r>
      <w:r><w:t>[1]</w:t></w:r>
      <w:r><w:fldChar w:fldCharType="end"/></w:r>
    </w:p>
  </w:body>
</w:document>
"""
    with zipfile.ZipFile(input_docx, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/media/image1.png", b"word-rewritten-large-media")
    with zipfile.ZipFile(media_source_docx, "w") as archive:
        archive.writestr("word/media/image1.png", b"original-small-media")

    optimize_docx_media(input_docx, media_source_docx, output_docx)
    audit = audit_zotero_docx(output_docx, tmp_path)
    with zipfile.ZipFile(output_docx) as archive:
        restored_media = archive.read("word/media/image1.png")
        restored_xml = archive.read("word/document.xml").decode("utf-8")

    assert restored_media == b"original-small-media"
    assert "ADDIN ZOTERO_ITEM CSL_CITATION" in restored_xml
    assert audit["csl_citation_fields"] == 1
