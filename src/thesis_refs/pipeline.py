from __future__ import annotations

import json
import io
import os
import re
import shutil
import tomllib
import uuid
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from openpyxl import Workbook

from thesis_refs.zotero_metadata import (
    creators_from_responsibility,
    creators_to_bibtex,
    creators_to_csl,
    is_corporate_author as zotero_is_corporate_author,
    split_author_parts as zotero_split_author_parts,
)


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT_DOC = ROOT / "examples" / "demo" / "raw" / "thesis.docx"
INPUT_DOC_ENV = "THESIS_REFS_INPUT_DOCX"
OUTPUT_ROOT_ENV = "THESIS_REFS_OUTPUT_ROOT"


def resolve_input_docx(input_docx: Path | str | None = None) -> Path:
    return Path(input_docx or os.environ.get(INPUT_DOC_ENV, DEFAULT_INPUT_DOC))


def default_output_dir(input_docx: Path) -> Path:
    try:
        is_demo_input = input_docx.resolve() == DEFAULT_INPUT_DOC.resolve()
    except FileNotFoundError:
        is_demo_input = input_docx == DEFAULT_INPUT_DOC
    if is_demo_input:
        return ROOT / "examples" / "demo" / "output"
    return ROOT / "output"


def unified_doc_name(input_docx: Path) -> str:
    try:
        is_demo_input = input_docx.resolve() == DEFAULT_INPUT_DOC.resolve()
    except FileNotFoundError:
        is_demo_input = input_docx == DEFAULT_INPUT_DOC
    return "论文_统一编号_demo.docx" if is_demo_input else "论文_统一编号.docx"


INPUT_DOC = resolve_input_docx()
OUTPUT_DIR = Path(os.environ.get(OUTPUT_ROOT_ENV, default_output_dir(INPUT_DOC)))
BUILD_DIR = OUTPUT_DIR / "build"
UNIFIED_DOC = OUTPUT_DIR / "doc" / unified_doc_name(INPUT_DOC)


def configure_paths(input_docx: Path | str | None = None, output_root: Path | str | None = None) -> Path:
    """Set process-local input/output defaults for CLI scripts."""
    global INPUT_DOC, OUTPUT_DIR, BUILD_DIR, UNIFIED_DOC

    INPUT_DOC = resolve_input_docx(input_docx)
    OUTPUT_DIR = Path(output_root or os.environ.get(OUTPUT_ROOT_ENV, default_output_dir(INPUT_DOC)))
    BUILD_DIR = OUTPUT_DIR / "build"
    UNIFIED_DOC = OUTPUT_DIR / "doc" / unified_doc_name(INPUT_DOC)
    return INPUT_DOC


def load_pipeline_config(config_path: Path | str | None = None) -> dict[str, object]:
    if not config_path:
        return {}
    path = Path(config_path)
    if not path.exists():
        raise FileNotFoundError(f"Config not found: {path}")
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
    else:
        data = tomllib.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Pipeline config must be a table/object.")
    allowed = {"input_docx", "output_root", "zotero_group_id"}
    unknown = sorted(set(data) - allowed)
    if unknown:
        raise ValueError(f"Unsupported config keys: {', '.join(unknown)}")
    return data

CITATION_RE = re.compile(r"\[(\d{1,3})\]")
CITATION_CLUSTER_RE = re.compile(r"(?:\[\s*\d{1,3}(?:\s*[-–]\s*\d{1,3})?(?:\s*,\s*\d{1,3}(?:\s*[-–]\s*\d{1,3})?)*\s*\]\s*)+")
PANDOC_CITATION_RE = re.compile(r"\[(?:-?@[A-Za-z0-9_:.#/-]+(?:\s*;\s*-?@[A-Za-z0-9_:.#/-]+)*)\]")
PANDOC_CITATION_KEY_RE = re.compile(r"-?@([A-Za-z0-9_:.#/-]+)")
OPERATOR_PLACEHOLDER_RE = re.compile(r"\[\[ZOTERO_P\d{3}\]\]")
ZOTERO_BIBLIOGRAPHY_ANCHOR_RE = re.compile(r"\[\[ZOTERO_BIBLIOGRAPHY(?:\]\])?")
DOI_RESOLVER_RE = re.compile(
    r"https?://(?:dx\.)?doi\.org/(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)",
    flags=re.IGNORECASE,
)
URL_TEXT_RE = re.compile(r"\s+https?://[^\s<]+", flags=re.IGNORECASE)
DOI_TEXT_RE = re.compile(r"\s+DOI:\s*10\.\d{4,9}/[-._;()/:A-Za-z0-9]+\.?", flags=re.IGNORECASE)
ALL_CAPS_ENGLISH_AUTHOR_RE = re.compile(r"\b[A-Z]{2,}\s+[A-Z](?:\s+[A-Z])?\b")
LATEX_SUPERSCRIPT_RE = re.compile(r"([A-Za-z])\$\^\{([0-9]+)\}\$")
SUPERSCRIPT_DIGITS = str.maketrans("0123456789", "⁰¹²³⁴⁵⁶⁷⁸⁹")
BIBLIOGRAPHY_HANGING_TWIPS = "480"
BIBLIOGRAPHY_LINE_SPACING = "300"
REF_PREFIX_RE = re.compile(r"^\[(\d{1,3})\]\s*(.+)")
YEAR_RE = re.compile(r"(19|20)\d{2}")
DOI_RE = re.compile(r"10\.\d{4,9}/[-._;()/:A-Za-z0-9]+")
URL_RE = re.compile(r"https?://\S+")
TYPE_RE = re.compile(r"\[(EB/OL|R/OL|J|C|S/OL|S|M|D|N|P|OL)\]")
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
ElementTree.register_namespace("w", WORD_NS)


@dataclass
class Reference:
    ref_id: str
    source_chapter: str
    source_section: str
    local_id: int
    raw_text: str
    paragraph_index: int
    ref_type: str
    dedupe_key: str
    title: str
    responsibility: str
    year: str | None = None
    doi: str | None = None
    url: str | None = None
    global_id: int | None = None
    verification_status: str = "manual_review_needed"
    verification_evidence: str = ""
    verification_note: str = ""


@dataclass(frozen=True)
class ReferenceOverride:
    contains: str
    raw_text: str | None = None
    ref_type: str | None = None
    year: str | None = None
    url: str | None = None
    status: str = "manual_review_needed"
    evidence: str = ""
    note: str = ""


REFERENCE_OVERRIDES: list[ReferenceOverride] = []

CHINESE_VERIFICATION_EVIDENCE: dict[str, tuple[str, str, str]] = {}


@dataclass
class CitationOccurrence:
    paragraph_index: int
    marker_index: int
    source_chapter: str
    source_section: str
    local_id: int
    ref_id: str


@dataclass
class ReferenceModel:
    references: list[Reference]
    citation_occurrences: list[CitationOccurrence]
    global_id_by_ref_id: dict[str, int]
    unique_references: list[Reference]
    input_docx_path: str | None = None
    reference_paragraphs: list[int] = field(default_factory=list)
    missing_reference_blocks: list[dict[str, object]] = field(default_factory=list)
    chapter_one_reference_paragraphs: list[int] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "references": [asdict(ref) for ref in self.references],
            "citation_occurrences": [asdict(occ) for occ in self.citation_occurrences],
            "global_id_by_ref_id": self.global_id_by_ref_id,
            "unique_references": [asdict(ref) for ref in self.unique_references],
            "input_docx_path": self.input_docx_path,
            "reference_paragraphs": self.reference_paragraphs,
            "missing_reference_blocks": self.missing_reference_blocks,
            "chapter_one_reference_paragraphs": self.chapter_one_reference_paragraphs,
        }


def normalize_space(text: str) -> str:
    return " ".join((text or "").split())


def display_path(path: Path | str) -> str:
    path = Path(path)
    try:
        return str(path.resolve().relative_to(ROOT.resolve()))
    except (FileNotFoundError, ValueError):
        return str(path)


def strip_reference_prefix(text: str) -> tuple[int | None, str]:
    clean = normalize_space(text)
    match = REF_PREFIX_RE.match(clean)
    if not match:
        return None, clean
    return int(match.group(1)), normalize_space(match.group(2))


def reference_type(text: str) -> str:
    match = TYPE_RE.search(text)
    return match.group(1) if match else "misc"


def extract_year(text: str) -> str | None:
    full_matches = re.findall(r"(?:19|20)\d{2}", text)
    return full_matches[0] if full_matches else None


def extract_doi(text: str) -> str | None:
    match = DOI_RE.search(text)
    return match.group(0).rstrip(".") if match else None


def extract_url(text: str) -> str | None:
    match = URL_RE.search(text)
    return match.group(0).rstrip(".") if match else None


def split_responsibility_and_title(text: str) -> tuple[str, str]:
    clean = strip_reference_prefix(text)[1]
    marker = TYPE_RE.search(clean)
    before_type = clean[: marker.start()] if marker else clean
    if "." in before_type:
        responsibility, title = before_type.split(".", 1)
        return normalize_space(responsibility), normalize_space(title)
    if "．" in before_type:
        responsibility, title = before_type.split("．", 1)
        return normalize_space(responsibility), normalize_space(title)
    return "", normalize_space(before_type)


def normalize_key_part(text: str) -> str:
    text = text.lower()
    text = re.sub(r"\bet\s+al\b", "", text)
    text = re.sub(r"[《》“”\"'`·、，,。．.\s:：;；\-—_()（）\[\]/]+", "", text)
    return text


def make_dedupe_key(raw_text: str) -> str:
    responsibility, title = split_responsibility_and_title(raw_text)
    title_key = normalize_key_part(title)
    responsibility_key = normalize_key_part(responsibility)[:18]
    year = extract_year(raw_text) or ""
    if title_key:
        return f"{title_key}|{responsibility_key}|{year}"
    return normalize_key_part(strip_reference_prefix(raw_text)[1])


def make_reference(
    source_chapter: str,
    source_section: str,
    local_id: int,
    raw_text: str,
    paragraph_index: int,
) -> Reference:
    body = strip_reference_prefix(raw_text)[1]
    responsibility, title = split_responsibility_and_title(body)
    ref_id = f"{source_chapter}:{source_section}:{local_id}"
    ref = Reference(
        ref_id=ref_id,
        source_chapter=source_chapter,
        source_section=source_section,
        local_id=local_id,
        raw_text=body,
        paragraph_index=paragraph_index,
        ref_type=reference_type(body),
        dedupe_key=make_dedupe_key(body),
        title=title,
        responsibility=responsibility,
        year=extract_year(body),
        doi=extract_doi(body),
        url=extract_url(body),
    )
    apply_reference_overrides(ref)
    return ref


def apply_reference_overrides(ref: Reference) -> None:
    for override in REFERENCE_OVERRIDES:
        if override.contains not in ref.raw_text:
            continue
        if override.raw_text:
            ref.raw_text = override.raw_text
            ref.responsibility, ref.title = split_responsibility_and_title(ref.raw_text)
            ref.dedupe_key = make_dedupe_key(ref.raw_text)
            ref.doi = extract_doi(ref.raw_text)
        if override.ref_type:
            ref.ref_type = override.ref_type
        ref.year = override.year
        ref.url = override.url or ref.url
        ref.verification_status = override.status
        ref.verification_evidence = override.evidence
        ref.verification_note = override.note
        return

    if ref.source_chapter == "第一章":
        for needle, (status, evidence, note) in CHINESE_VERIFICATION_EVIDENCE.items():
            if needle in ref.raw_text:
                ref.verification_status = status
                ref.verification_evidence = evidence
                ref.verification_note = note
                return


def paragraph_texts(doc: DocumentObject) -> list[str]:
    return [normalize_space(paragraph.text) for paragraph in doc.paragraphs]


def is_heading(paragraph) -> bool:
    return paragraph.style.name.startswith("Heading")


def find_heading_index(doc: DocumentObject, prefix: str, start: int = 0) -> int:
    for idx, paragraph in enumerate(doc.paragraphs[start:], start):
        if paragraph.style.name == "Heading 1" and normalize_space(paragraph.text).startswith(prefix):
            return idx
    raise ValueError(f"Heading not found: {prefix}")


def chapter_range(doc: DocumentObject, heading_prefix: str) -> tuple[int, int]:
    start = find_heading_index(doc, heading_prefix)
    end = len(doc.paragraphs)
    for idx, paragraph in enumerate(doc.paragraphs[start + 1 :], start + 1):
        if paragraph.style.name == "Heading 1":
            end = idx
            break
    return start, end


def section_ranges(doc: DocumentObject, start: int, end: int) -> list[tuple[str, int, int]]:
    starts: list[tuple[str, int]] = []
    for idx in range(start + 1, end):
        paragraph = doc.paragraphs[idx]
        text = normalize_space(paragraph.text)
        if paragraph.style.name in {"Heading 2", "Heading 3"} and text:
            starts.append((text, idx))
    ranges: list[tuple[str, int, int]] = []
    for pos, (title, idx) in enumerate(starts):
        next_idx = starts[pos + 1][1] if pos + 1 < len(starts) else end
        ranges.append((title, idx, next_idx))
    return ranges


def heading1_ranges(doc: DocumentObject) -> list[tuple[str, int, int]]:
    starts: list[tuple[str, int]] = []
    for idx, paragraph in enumerate(doc.paragraphs):
        text = normalize_space(paragraph.text)
        if paragraph.style.name == "Heading 1" and text:
            starts.append((text, idx))
    ranges: list[tuple[str, int, int]] = []
    for pos, (title, idx) in enumerate(starts):
        next_idx = starts[pos + 1][1] if pos + 1 < len(starts) else len(doc.paragraphs)
        ranges.append((title, idx, next_idx))
    return ranges


def final_bibliography_heading_index(doc: DocumentObject) -> int | None:
    matches = [
        idx
        for idx, paragraph in enumerate(doc.paragraphs)
        if paragraph.style.name == "Heading 1" and normalize_space(paragraph.text) == "参考文献"
    ]
    return matches[-1] if matches else None


def reference_blocks_in_range(doc: DocumentObject, start: int, end: int) -> list[list[tuple[int, int, str]]]:
    blocks: list[list[tuple[int, int, str]]] = []
    current: list[tuple[int, int, str]] = []
    for idx in range(start, end):
        paragraph = doc.paragraphs[idx]
        if is_heading(paragraph):
            if current:
                blocks.append(current)
                current = []
            continue
        text = normalize_space(paragraph.text)
        local_id, body = strip_reference_prefix(text)
        if local_id is None:
            if current:
                blocks.append(current)
                current = []
            continue
        current.append((idx, local_id, body))
    if current:
        blocks.append(current)
    return blocks


def nearest_section_before(doc: DocumentObject, chapter_title: str, chapter_start: int, before_idx: int) -> tuple[str, int]:
    section_title = chapter_title
    section_start = chapter_start
    for idx in range(chapter_start + 1, before_idx):
        paragraph = doc.paragraphs[idx]
        text = normalize_space(paragraph.text)
        if paragraph.style.name in {"Heading 2", "Heading 3"} and text:
            section_title = text
            section_start = idx
    return section_title, section_start


def extract_references_from_single_document(
    doc: DocumentObject,
) -> tuple[list[Reference], list[CitationOccurrence], list[int], list[dict[str, object]]]:
    refs: list[Reference] = []
    occurrences: list[CitationOccurrence] = []
    ref_paragraphs: list[int] = []
    missing_blocks: list[dict[str, object]] = []
    bibliography_idx = final_bibliography_heading_index(doc)

    for chapter_title, chapter_start, chapter_end in heading1_ranges(doc):
        if bibliography_idx is not None and chapter_start >= bibliography_idx:
            break
        if chapter_title in {"摘要", "Abstract", "目录"}:
            continue
        effective_end = min(chapter_end, bibliography_idx) if bibliography_idx is not None else chapter_end
        if effective_end <= chapter_start:
            continue

        chapter_blocks = reference_blocks_in_range(doc, chapter_start + 1, effective_end)
        if not chapter_blocks:
            citation_markers = [
                match.group(0)
                for idx in range(chapter_start + 1, effective_end)
                if not is_heading(doc.paragraphs[idx])
                for match in CITATION_RE.finditer(normalize_space(doc.paragraphs[idx].text))
            ]
            if citation_markers:
                missing_blocks.append(
                    {
                        "source_chapter": chapter_title,
                        "paragraph_range": [chapter_start + 1, effective_end],
                        "citation_markers": citation_markers[:20],
                        "issue": "missing_reference_block",
                    }
                )
            continue

        for block in chapter_blocks:
            block_start = block[0][0]
            section_title, section_start = nearest_section_before(doc, chapter_title, chapter_start, block_start)
            local_refs: dict[int, Reference] = {}
            for idx, local_id, body in block:
                ref = make_reference(chapter_title, section_title, local_id, body, idx)
                local_refs[local_id] = ref
                refs.append(ref)
                ref_paragraphs.append(idx)

            for idx in range(section_start + 1, block_start):
                paragraph = doc.paragraphs[idx]
                if is_heading(paragraph):
                    continue
                text = normalize_space(paragraph.text)
                if not text:
                    continue
                for marker_idx, match in enumerate(CITATION_RE.finditer(text)):
                    local_id = int(match.group(1))
                    ref = local_refs.get(local_id)
                    if ref is None:
                        continue
                    occurrences.append(
                        CitationOccurrence(
                            paragraph_index=idx,
                            marker_index=marker_idx,
                            source_chapter=chapter_title,
                            source_section=section_title,
                            local_id=local_id,
                            ref_id=ref.ref_id,
                        )
                    )

    return refs, occurrences, ref_paragraphs, missing_blocks


def extract_chapter_one(main_doc: DocumentObject) -> tuple[list[Reference], list[CitationOccurrence], list[int]]:
    refs: list[Reference] = []
    occurrences: list[CitationOccurrence] = []
    ref_paragraphs: list[int] = []
    chapter_start, chapter_end = chapter_range(main_doc, "第一章")

    for section_title, section_start, section_end in section_ranges(main_doc, chapter_start, chapter_end):
        local_refs: dict[int, Reference] = {}
        first_ref_idx: int | None = None
        for idx in range(section_start + 1, section_end):
            text = normalize_space(main_doc.paragraphs[idx].text)
            local_id, body = strip_reference_prefix(text)
            if local_id is None:
                continue
            if first_ref_idx is None:
                first_ref_idx = idx
            ref = make_reference("第一章", section_title, local_id, body, idx)
            local_refs[local_id] = ref
            refs.append(ref)
            ref_paragraphs.append(idx)

        if not local_refs:
            continue

        body_end = first_ref_idx if first_ref_idx is not None else section_end
        for idx in range(section_start + 1, body_end):
            text = normalize_space(main_doc.paragraphs[idx].text)
            for marker_idx, match in enumerate(CITATION_RE.finditer(text)):
                local_id = int(match.group(1))
                if local_id not in local_refs:
                    continue
                occurrences.append(
                    CitationOccurrence(
                        paragraph_index=idx,
                        marker_index=marker_idx,
                        source_chapter="第一章",
                        source_section=section_title,
                        local_id=local_id,
                        ref_id=local_refs[local_id].ref_id,
                    )
                )
    return refs, occurrences, ref_paragraphs


def assign_global_numbers(
    references: list[Reference],
    occurrences: list[CitationOccurrence],
) -> tuple[dict[str, int], list[Reference]]:
    ref_by_id = {ref.ref_id: ref for ref in references}
    key_to_global_id: dict[str, int] = {}
    global_id_by_ref_id: dict[str, int] = {}
    unique_refs: list[Reference] = []

    for occurrence in sorted(occurrences, key=lambda occ: (occ.paragraph_index, occ.marker_index)):
        ref = ref_by_id[occurrence.ref_id]
        if ref.dedupe_key not in key_to_global_id:
            global_id = len(unique_refs) + 1
            key_to_global_id[ref.dedupe_key] = global_id
            ref.global_id = global_id
            unique_refs.append(ref)
        global_id_by_ref_id[ref.ref_id] = key_to_global_id[ref.dedupe_key]

    for ref in references:
        if ref.ref_id in global_id_by_ref_id:
            continue
        if ref.dedupe_key not in key_to_global_id:
            global_id = len(unique_refs) + 1
            key_to_global_id[ref.dedupe_key] = global_id
            ref.global_id = global_id
            unique_refs.append(ref)
        global_id_by_ref_id[ref.ref_id] = key_to_global_id[ref.dedupe_key]

    for ref in references:
        ref.global_id = global_id_by_ref_id[ref.ref_id]
    for unique_ref in unique_refs:
        unique_ref.global_id = key_to_global_id[unique_ref.dedupe_key]

    return global_id_by_ref_id, unique_refs


def build_reference_model(input_docx_path: Path | str | None = None) -> ReferenceModel:
    input_path = resolve_input_docx(input_docx_path)
    doc = Document(input_path)
    references, occurrences, ref_paragraphs, missing_blocks = extract_references_from_single_document(doc)
    global_map, unique_refs = assign_global_numbers(references, occurrences)
    return ReferenceModel(
        references=references,
        citation_occurrences=occurrences,
        global_id_by_ref_id=global_map,
        unique_references=unique_refs,
        input_docx_path=str(input_path),
        reference_paragraphs=ref_paragraphs,
        missing_reference_blocks=missing_blocks,
        chapter_one_reference_paragraphs=ref_paragraphs,
    )


def compress_citation_runs(text: str) -> str:
    def replace_run(match: re.Match[str]) -> str:
        numbers = [int(item) for item in re.findall(r"\[(\d{1,3})\]", match.group(0))]
        if len(numbers) < 2:
            return match.group(0)
        parts: list[str] = []
        start = prev = numbers[0]
        for number in numbers[1:]:
            if number == prev + 1:
                prev = number
                continue
            parts.append(f"[{start}-{prev}]" if start != prev else f"[{start}]")
            start = prev = number
        parts.append(f"[{start}-{prev}]" if start != prev else f"[{start}]")
        return "".join(parts)

    return re.sub(r"(?:\[\d{1,3}\]\s*){2,}", replace_run, text)


def expand_citation_cluster(marker_text: str) -> list[int]:
    numbers: list[int] = []
    for content in re.findall(r"\[([^\]]+)\]", marker_text):
        for part in re.split(r"\s*,\s*", content):
            clean = normalize_space(part)
            if not clean:
                continue
            range_match = re.fullmatch(r"(\d{1,3})\s*[-–]\s*(\d{1,3})", clean)
            if range_match:
                start, end = int(range_match.group(1)), int(range_match.group(2))
                step = 1 if start <= end else -1
                numbers.extend(range(start, end + step, step))
                continue
            if re.fullmatch(r"\d{1,3}", clean):
                numbers.append(int(clean))
    return numbers


def _chapter_body_static_paragraphs(target_docx_path: Path) -> list[dict[str, object]]:
    doc = Document(target_docx_path)
    rows: list[dict[str, object]] = []
    bibliography_idx = final_bibliography_heading_index(doc)
    body_end = bibliography_idx if bibliography_idx is not None else len(doc.paragraphs)
    current_chapter = ""
    current_section = ""

    for idx, paragraph in enumerate(doc.paragraphs[:body_end]):
        text = normalize_space(paragraph.text)
        if paragraph.style.name == "Heading 1":
            current_chapter = text
            current_section = text
            continue
        if paragraph.style.name in {"Heading 2", "Heading 3"} and text:
            current_section = text
            continue
        if is_heading(paragraph) or "[" not in paragraph.text:
            continue
        if CITATION_CLUSTER_RE.search(text):
            rows.append(
                {
                    "paragraph_index": idx,
                    "source_chapter": current_chapter,
                    "source_section": current_section or current_chapter,
                    "paragraph_text": normalize_space(text),
                }
            )

    return rows


def generate_zotero_migration_checklist(
    model: ReferenceModel,
    target_docx_path: Path = UNIFIED_DOC,
) -> dict[str, object]:
    if not Path(target_docx_path).exists():
        raise FileNotFoundError(f"Unified DOCX not found: {target_docx_path}")
    titles_by_id = {
        int(ref.global_id or 0): ref.title or ref.raw_text
        for ref in model.unique_references
        if ref.global_id is not None
    }
    positions: list[dict[str, object]] = []
    for paragraph in _chapter_body_static_paragraphs(Path(target_docx_path)):
        citation_index = 0
        paragraph_text = str(paragraph["paragraph_text"])
        for match in CITATION_CLUSTER_RE.finditer(paragraph_text):
            global_ids = expand_citation_cluster(match.group(0))
            if not global_ids:
                continue
            citation_index += 1
            position_id = f"P{len(positions) + 1:03d}"
            positions.append(
                {
                    "position_id": position_id,
                    "paragraph_index": paragraph["paragraph_index"],
                    "citation_index_in_paragraph": citation_index,
                    "source_chapter": paragraph["source_chapter"],
                    "source_section": paragraph["source_section"],
                    "marker_text": normalize_space(match.group(0)),
                    "global_ids": global_ids,
                    "zotero_keys": [f"ref{number:03d}" for number in global_ids],
                    "titles": [titles_by_id.get(number, "") for number in global_ids],
                    "paragraph_text": paragraph_text,
                    "search_hint": paragraph_text[max(0, match.start() - 30): match.end() + 30],
                }
            )

    covered = sum(len(position["global_ids"]) for position in positions)
    expected_sequence: list[int] = []
    for occurrence in sorted(model.citation_occurrences, key=lambda item: (item.paragraph_index, item.marker_index)):
        expected_sequence.append(model.global_id_by_ref_id[occurrence.ref_id])
    generated_sequence: list[int] = []
    for position in positions:
        generated_sequence.extend(int(item) for item in position["global_ids"])
    mismatches = []
    if expected_sequence != generated_sequence:
        mismatch_at = next(
            (
                index
                for index, (expected, generated) in enumerate(zip(expected_sequence, generated_sequence), start=1)
                if expected != generated
            ),
            min(len(expected_sequence), len(generated_sequence)) + 1,
        )
        mismatches.append(
            {
                "occurrence_index": mismatch_at,
                "expected_global_ids": expected_sequence[max(0, mismatch_at - 4): mismatch_at + 3],
                "generated_global_ids": generated_sequence[max(0, mismatch_at - 4): mismatch_at + 3],
            }
        )
    invalid_global_ids = sorted(
        {
            global_id
            for global_id in generated_sequence
            if global_id < 1 or global_id > len(model.unique_references)
        }
    )
    return {
        "summary": {
            "target_docx": display_path(target_docx_path),
            "unique_references": len(model.unique_references),
            "citation_positions": len(positions),
            "citation_paragraphs": len({position["paragraph_index"] for position in positions}),
            "covered_citation_occurrences": covered,
            "expected_citation_occurrences": len(model.citation_occurrences),
            "paragraph_coverage_ok": not mismatches and not invalid_global_ids,
            "paragraph_coverage_mismatches": len(mismatches),
            "invalid_global_ids": invalid_global_ids,
        },
        "coverage_mismatches": mismatches,
        "positions": positions,
    }


def write_zotero_migration_checklist(
    model: ReferenceModel,
    output_dir: Path | None = None,
    target_docx_path: Path | None = None,
) -> tuple[Path, Path]:
    output_dir = Path(output_dir or OUTPUT_DIR / "zotero")
    target_docx_path = Path(target_docx_path or UNIFIED_DOC)
    output_dir.mkdir(parents=True, exist_ok=True)
    if not target_docx_path.exists():
        rewrite_main_docx(model, target_docx_path, Path(model.input_docx_path or INPUT_DOC))
    checklist = generate_zotero_migration_checklist(model, target_docx_path)
    json_path = output_dir / "migration_checklist.json"
    md_path = output_dir / "migration_checklist.md"
    write_json(json_path, checklist)

    summary = checklist["summary"]
    lines = [
        "# Zotero 动态域迁移清单",
        "",
        f"- Target DOCX: `{summary['target_docx']}`",
        f"- Unique references: {summary['unique_references']}",
        f"- Citation positions: {summary['citation_positions']}",
        f"- Citation paragraphs: {summary['citation_paragraphs']}",
        f"- Covered citation occurrences: {summary['covered_citation_occurrences']} / {summary['expected_citation_occurrences']}",
        f"- Paragraph coverage OK: {summary['paragraph_coverage_ok']}",
        "",
        "## 操作规则",
        "",
        "- 在 Word 中删除 `Marker` 对应的静态文本，再用 Zotero Add/Edit Citation 插入 `Zotero keys` 中的条目。",
        "- 同一行包含多个 key 时，在同一个 Zotero citation 中连续添加多篇文献。",
        "- 每处理 5-10 个位置保存一个 checkpoint 副本。",
        "",
        "| Position | Paragraph | Section | Marker | Zotero keys | Titles |",
        "|---|---:|---|---|---|---|",
    ]
    for position in checklist["positions"]:
        titles = "; ".join(str(title).replace("|", "/") for title in position["titles"])
        lines.append(
            f"| {position['position_id']} | {position['paragraph_index']} | "
            f"{position['source_section']} | `{position['marker_text']}` | "
            f"`{', '.join(position['zotero_keys'])}` | {titles} |"
        )
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return json_path, md_path


def pandoc_marker_for_keys(zotero_keys: list[str]) -> str:
    return "[" + "; ".join(f"@{key}" for key in zotero_keys) + "]"


def operator_placeholder_for_position(position_id: str) -> str:
    return f"[[ZOTERO_{position_id}]]"


def _text_nodes_for_paragraph(paragraph: Paragraph) -> list[ElementTree.Element]:
    return list(paragraph._p.iter(qn("w:t")))


def _replace_text_span_in_nodes(
    nodes: list[ElementTree.Element],
    start: int,
    end: int,
    replacement: str,
) -> None:
    segments: list[tuple[ElementTree.Element, str, int, int]] = []
    cursor = 0
    for node in nodes:
        text = node.text or ""
        next_cursor = cursor + len(text)
        segments.append((node, text, cursor, next_cursor))
        cursor = next_cursor

    overlapping = [
        (node, text, seg_start, seg_end)
        for node, text, seg_start, seg_end in segments
        if seg_start < end and seg_end > start
    ]
    if not overlapping:
        raise ValueError("No text node overlaps the requested replacement span.")

    if len(overlapping) == 1:
        node, text, seg_start, _ = overlapping[0]
        local_start = start - seg_start
        local_end = end - seg_start
        node.text = text[:local_start] + replacement + text[local_end:]
        return

    first_node, first_text, first_start, _ = overlapping[0]
    last_node, last_text, last_start, _ = overlapping[-1]
    first_node.text = first_text[: start - first_start] + replacement
    for node, _, _, _ in overlapping[1:-1]:
        node.text = ""
    last_node.text = last_text[end - last_start :]


def replace_nth_citation_cluster(
    paragraph: Paragraph,
    citation_index_in_paragraph: int,
    replacement: str,
    expected_marker: str | None = None,
) -> str:
    nodes = _text_nodes_for_paragraph(paragraph)
    full_text = "".join(node.text or "" for node in nodes)
    matches = list(CITATION_CLUSTER_RE.finditer(full_text))
    if citation_index_in_paragraph < 1 or citation_index_in_paragraph > len(matches):
        raise ValueError(
            f"Citation cluster #{citation_index_in_paragraph} not found in paragraph: "
            f"{normalize_space(full_text)[:120]}"
        )
    match = matches[citation_index_in_paragraph - 1]
    marker = normalize_space(match.group(0))
    if expected_marker is not None and marker != normalize_space(expected_marker):
        raise ValueError(f"Expected marker {expected_marker!r}, found {marker!r}.")
    _replace_text_span_in_nodes(nodes, match.start(), match.end(), replacement)
    return marker


def _write_operator_checklist(
    checklist: dict[str, object],
    json_path: Path,
    md_path: Path,
    mode: str,
    target_docx_path: Path,
) -> None:
    json_path.parent.mkdir(parents=True, exist_ok=True)
    summary = dict(checklist["summary"])
    summary["target_docx"] = display_path(target_docx_path)
    summary["migration_mode"] = mode
    output = {
        **checklist,
        "summary": summary,
    }
    write_json(json_path, output)

    title = "Zotero Scan 输入清单" if mode == "pandoc_scan" else "Zotero Operator 占位符清单"
    lines = [
        f"# {title}",
        "",
        f"- Target DOCX: `{target_docx_path}`",
        f"- Migration mode: `{mode}`",
        f"- Citation positions: {summary['citation_positions']}",
        f"- Covered citation occurrences: {summary['covered_citation_occurrences']} / {summary['expected_citation_occurrences']}",
        "",
        "| Position | Paragraph | Original Marker | Replacement | Zotero keys | Titles |",
        "|---|---:|---|---|---|---|",
    ]
    for position in checklist["positions"]:
        zotero_keys = list(position["zotero_keys"])
        if mode == "pandoc_scan":
            replacement = pandoc_marker_for_keys(zotero_keys)
        else:
            replacement = operator_placeholder_for_position(str(position["position_id"]))
        titles = "; ".join(str(title).replace("|", "/") for title in position["titles"])
        lines.append(
            f"| {position['position_id']} | {position['paragraph_index']} | "
            f"`{position['marker_text']}` | `{replacement}` | "
            f"`{', '.join(zotero_keys)}` | {titles} |"
        )
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _remove_final_static_references_and_insert_anchor(
    doc: DocumentObject,
    anchor_text: str = "[[ZOTERO_BIBLIOGRAPHY]]",
) -> None:
    ref_heading_idx = find_heading_index(doc, "参考文献")
    insert_idx = ref_heading_idx + 1
    while insert_idx < len(doc.paragraphs):
        paragraph = doc.paragraphs[insert_idx]
        if paragraph.style.name == "Heading 1":
            break
        delete_paragraph(paragraph)
    insert_paragraph_after(doc.paragraphs[ref_heading_idx], anchor_text, style="Normal")


def _prepare_zotero_replacement_docx(
    model: ReferenceModel,
    source_docx_path: Path,
    output_docx_path: Path,
    output_dir: Path,
    mode: str,
) -> tuple[Path, Path, Path]:
    source_docx_path = Path(source_docx_path)
    output_docx_path = Path(output_docx_path)
    output_dir = Path(output_dir)
    if not source_docx_path.exists():
        rewrite_main_docx(model, source_docx_path, Path(model.input_docx_path or INPUT_DOC))
    checklist = generate_zotero_migration_checklist(model, source_docx_path)
    if not checklist["summary"]["paragraph_coverage_ok"]:
        raise ValueError("Migration checklist coverage is not OK; refusing to prepare DOCX.")

    output_docx_path.parent.mkdir(parents=True, exist_ok=True)
    working_path = output_docx_path.parent / f".{output_docx_path.stem}.{uuid.uuid4().hex}.tmp.docx"
    shutil.copyfile(source_docx_path, working_path)
    doc = Document(working_path)

    positions_by_paragraph: dict[int, list[dict[str, object]]] = {}
    for position in checklist["positions"]:
        positions_by_paragraph.setdefault(int(position["paragraph_index"]), []).append(position)

    for paragraph_index, positions in positions_by_paragraph.items():
        paragraph = doc.paragraphs[paragraph_index]
        for position in sorted(positions, key=lambda item: int(item["citation_index_in_paragraph"]), reverse=True):
            zotero_keys = list(position["zotero_keys"])
            if mode == "pandoc_scan":
                replacement = pandoc_marker_for_keys(zotero_keys)
            elif mode == "operator_placeholders":
                replacement = operator_placeholder_for_position(str(position["position_id"]))
            else:
                raise ValueError(f"Unsupported Zotero replacement mode: {mode}")
            replace_nth_citation_cluster(
                paragraph,
                int(position["citation_index_in_paragraph"]),
                replacement,
                str(position["marker_text"]),
            )

    _remove_final_static_references_and_insert_anchor(doc)
    doc.save(working_path)
    shutil.move(str(working_path), output_docx_path)

    if mode == "pandoc_scan":
        json_path = output_dir / "scan_migration_checklist.json"
        md_path = output_dir / "scan_migration_checklist.md"
    else:
        json_path = output_dir / "operator_migration_checklist.json"
        md_path = output_dir / "operator_migration_checklist.md"
    _write_operator_checklist(checklist, json_path, md_path, mode, output_docx_path)
    return output_docx_path, json_path, md_path


def prepare_zotero_operator_docx(
    model: ReferenceModel,
    source_docx_path: Path | None = None,
    output_docx_path: Path | None = None,
    output_dir: Path | None = None,
) -> tuple[Path, Path, Path]:
    return _prepare_zotero_replacement_docx(
        model,
        Path(source_docx_path or UNIFIED_DOC),
        Path(output_docx_path or OUTPUT_DIR / "doc" / "论文_zotero_dynamic_operator.docx"),
        Path(output_dir or OUTPUT_DIR / "zotero"),
        mode="operator_placeholders",
    )


def _iter_docx_xml(docx_path: Path) -> Iterable[tuple[str, str]]:
    with zipfile.ZipFile(docx_path) as archive:
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                yield name, archive.read(name).decode("utf-8", errors="ignore")


def _is_visible_story_xml(name: str) -> bool:
    basename = Path(name).name
    return (
        basename == "document.xml"
        or basename.startswith("header")
        or basename.startswith("footer")
        or basename in {"footnotes.xml", "endnotes.xml", "comments.xml"}
    )


def _field_instructions(xml_text: str) -> list[str]:
    try:
        root = ElementTree.fromstring(xml_text.encode("utf-8"))
    except ElementTree.ParseError:
        return []

    instructions: list[str] = []
    current_complex: list[str] | None = None

    for element in root.iter():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "fldSimple":
            instr = element.attrib.get(f"{{{WORD_NS}}}instr")
            if instr:
                instructions.append(normalize_space(instr))
            continue
        if tag == "fldChar":
            field_type = element.attrib.get(f"{{{WORD_NS}}}fldCharType")
            if field_type == "begin":
                current_complex = []
            elif field_type == "separate" and current_complex is not None:
                if current_complex:
                    instructions.append(normalize_space("".join(current_complex)))
                current_complex = None
            elif field_type == "end" and current_complex is not None:
                instructions.append(normalize_space("".join(current_complex)))
                current_complex = None
            continue
        if tag == "instrText" and current_complex is not None and element.text:
            current_complex.append(element.text)
        elif tag == "instrText" and element.text:
            instructions.append(normalize_space(element.text))

    return [instruction for instruction in instructions if instruction]


def _visible_paragraphs_outside_fields(xml_text: str) -> list[str]:
    try:
        root = ElementTree.fromstring(xml_text.encode("utf-8"))
    except ElementTree.ParseError:
        return []

    paragraphs: list[str] = []
    chunks: list[str] = []
    simple_depth = 0
    complex_depth = 0

    def walk(element: ElementTree.Element) -> None:
        nonlocal simple_depth, complex_depth, chunks
        tag = element.tag.rsplit("}", 1)[-1]

        if tag == "fldSimple":
            simple_depth += 1
            for child in element:
                walk(child)
            simple_depth -= 1
            return

        if tag == "fldChar":
            field_type = element.attrib.get(f"{{{WORD_NS}}}fldCharType")
            if field_type == "begin":
                complex_depth += 1
            elif field_type == "end":
                complex_depth = max(0, complex_depth - 1)
            return

        if tag == "p":
            previous_chunks = chunks
            chunks = []
            for child in element:
                walk(child)
            text = "".join(chunks)
            if text:
                paragraphs.append(text)
            chunks = previous_chunks
            return

        if tag == "t" and simple_depth == 0 and complex_depth == 0 and element.text:
            chunks.append(element.text)
        for child in element:
            walk(child)

    walk(root)
    return paragraphs


def _field_structure_diagnostics(xml_text: str) -> dict[str, int]:
    try:
        root = ElementTree.fromstring(xml_text.encode("utf-8"))
    except ElementTree.ParseError:
        return {
            "complex_field_begin_count": 0,
            "complex_field_separate_count": 0,
            "complex_field_end_count": 0,
            "complex_field_unclosed_count": 0,
            "complex_field_unmatched_end_count": 0,
            "bibliography_field_unclosed_count": 0,
            "visible_field_code_text_count": 0,
        }

    stack: list[dict[str, object]] = []
    begin_count = 0
    separate_count = 0
    end_count = 0
    unmatched_end_count = 0
    visible_field_code_text_count = 0

    for element in root.iter():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "fldChar":
            field_type = element.attrib.get(f"{{{WORD_NS}}}fldCharType")
            if field_type == "begin":
                begin_count += 1
                stack.append({"instruction": [], "is_bibliography": False})
            elif field_type == "separate":
                separate_count += 1
                if stack:
                    instruction = normalize_space("".join(stack[-1]["instruction"]))
                    if "CSL_BIBLIOGRAPHY" in instruction:
                        stack[-1]["is_bibliography"] = True
            elif field_type == "end":
                end_count += 1
                if stack:
                    stack.pop()
                else:
                    unmatched_end_count += 1
            continue
        if tag == "instrText" and stack and element.text:
            stack[-1]["instruction"].append(element.text)
            continue
        if tag == "t" and element.text and re.search(r"ADDIN ZOTERO_|CSL_(?:CITATION|BIBLIOGRAPHY)", element.text):
            visible_field_code_text_count += 1

    bibliography_unclosed_count = 0
    for field in stack:
        instruction = normalize_space("".join(field["instruction"]))
        if field.get("is_bibliography") or "CSL_BIBLIOGRAPHY" in instruction:
            bibliography_unclosed_count += 1

    return {
        "complex_field_begin_count": begin_count,
        "complex_field_separate_count": separate_count,
        "complex_field_end_count": end_count,
        "complex_field_unclosed_count": len(stack),
        "complex_field_unmatched_end_count": unmatched_end_count,
        "bibliography_field_unclosed_count": bibliography_unclosed_count,
        "visible_field_code_text_count": visible_field_code_text_count,
    }


def _expected_zotero_key_sequence(checklist_path: Path | None) -> list[str]:
    if not checklist_path:
        return []
    try:
        checklist = json.loads(Path(checklist_path).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    sequence: list[str] = []
    for position in checklist.get("positions", []):
        sequence.extend(str(key) for key in position.get("zotero_keys", []))
    return sequence


def _read_checklist(checklist_path: Path | None) -> dict[str, object]:
    if not checklist_path:
        return {}
    try:
        checklist = json.loads(Path(checklist_path).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return checklist if isinstance(checklist, dict) else {}


def _expected_reference_types_by_number(checklist_path: Path | None) -> dict[int, str]:
    checklist = _read_checklist(checklist_path)
    expected: dict[int, str] = {}
    references = checklist.get("references")
    if isinstance(references, list):
        for item in references:
            if not isinstance(item, dict):
                continue
            number = item.get("global_id") or item.get("number")
            ref_type = item.get("ref_type") or item.get("type")
            try:
                number_int = int(number)
            except (TypeError, ValueError):
                continue
            if ref_type:
                expected[number_int] = str(ref_type)
    if expected:
        return expected

    try:
        model = build_reference_model()
    except Exception:
        return {}
    return {
        int(ref.global_id): ref.ref_type
        for ref in model.unique_references
        if ref.global_id is not None and ref.ref_type
    }


def _expected_operator_placeholders(checklist_path: Path | None) -> list[str]:
    if not checklist_path:
        return []
    try:
        checklist = json.loads(Path(checklist_path).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    return [
        operator_placeholder_for_position(str(position.get("position_id", "")))
        for position in checklist.get("positions", [])
    ]


def _reference_number_limit_from_keys(keys: list[str]) -> int | None:
    numbers: list[int] = []
    for key in keys:
        match = re.fullmatch(r"ref0*(\d+)", key)
        if match:
            numbers.append(int(match.group(1)))
    return max(numbers) if numbers else None


def _static_marker_within_reference_limit(marker: str, limit: int | None) -> bool:
    if limit is None:
        return True
    numbers = [int(item) for item in re.findall(r"\d{1,3}", marker)]
    return bool(numbers) and all(1 <= number <= limit for number in numbers)


def _pandoc_citation_items(paragraphs: list[str]) -> tuple[list[str], list[str]]:
    markers: list[str] = []
    keys: list[str] = []
    for paragraph in paragraphs:
        for marker in PANDOC_CITATION_RE.findall(paragraph):
            markers.append(marker)
            keys.extend(PANDOC_CITATION_KEY_RE.findall(marker))
    return markers, keys


def _operator_placeholders(paragraphs: list[str]) -> list[str]:
    return [
        placeholder
        for paragraph in paragraphs
        for placeholder in OPERATOR_PLACEHOLDER_RE.findall(paragraph)
    ]


def _json_payload_after_token(instruction: str, token: str) -> dict[str, object] | None:
    token_index = instruction.find(token)
    if token_index < 0:
        return None
    start = instruction.find("{", token_index + len(token))
    if start < 0:
        return None
    decoder = json.JSONDecoder()
    try:
        payload, _ = decoder.raw_decode(instruction[start:])
    except json.JSONDecodeError:
        return None
    return payload if isinstance(payload, dict) else None


def _citation_key_from_item(item: object) -> str | None:
    if not isinstance(item, dict):
        return None
    item_data = item.get("itemData")
    if isinstance(item_data, dict):
        citation_key = item_data.get("citation-key")
        if isinstance(citation_key, str) and citation_key:
            return citation_key
        note = item_data.get("note")
        if isinstance(note, str):
            match = re.search(r"citation-key:\s*([A-Za-z0-9_:.#/-]+)", note)
            if match:
                return match.group(1)
    return None


def _zotero_citation_items_from_instructions(
    instructions: list[str],
) -> tuple[int, list[str], int]:
    total_items = 0
    keys: list[str] = []
    parse_errors = 0
    for instruction in instructions:
        if "CSL_CITATION" not in instruction:
            continue
        payload = _json_payload_after_token(instruction, "CSL_CITATION")
        if payload is None:
            parse_errors += 1
            continue
        citation_items = payload.get("citationItems", [])
        if not isinstance(citation_items, list):
            parse_errors += 1
            continue
        total_items += len(citation_items)
        for item in citation_items:
            key = _citation_key_from_item(item)
            if key:
                keys.append(key)
    return total_items, keys, parse_errors


def _paragraph_text_with_tabs(paragraph: ElementTree.Element) -> str:
    chunks: list[str] = []
    for child in paragraph:
        if child.tag == f"{{{WORD_NS}}}pPr":
            continue
        for element in child.iter():
            tag = element.tag.rsplit("}", 1)[-1]
            if tag == "t" and element.text:
                chunks.append(element.text)
            elif tag == "tab":
                chunks.append("\t")
    return "".join(chunks)


def _child(parent: ElementTree.Element, tag: str) -> ElementTree.Element | None:
    return parent.find(f"w:{tag}", {"w": WORD_NS})


def _ensure_child(parent: ElementTree.Element, tag: str) -> ElementTree.Element:
    child = _child(parent, tag)
    if child is None:
        child = ElementTree.SubElement(parent, f"{{{WORD_NS}}}{tag}")
    return child


def _ensure_run_properties(run: ElementTree.Element) -> ElementTree.Element:
    rpr = _child(run, "rPr")
    if rpr is None:
        rpr = ElementTree.Element(f"{{{WORD_NS}}}rPr")
        run.insert(0, rpr)
    elif list(run).index(rpr) != 0:
        run.remove(rpr)
        run.insert(0, rpr)
    return rpr


def _style_id(paragraph: ElementTree.Element) -> str | None:
    ppr = _child(paragraph, "pPr")
    if ppr is None:
        return None
    pstyle = _child(ppr, "pStyle")
    if pstyle is None:
        return None
    return pstyle.attrib.get(f"{{{WORD_NS}}}val")


def _bibliography_paragraphs_from_document_xml(xml_text: str) -> list[tuple[ElementTree.Element, str]]:
    try:
        root = ElementTree.fromstring(xml_text.encode("utf-8"))
    except ElementTree.ParseError:
        return []

    paragraphs: list[tuple[ElementTree.Element, str]] = []
    in_references = False
    for paragraph in root.iter(f"{{{WORD_NS}}}p"):
        text = _paragraph_text_with_tabs(paragraph)
        stripped = normalize_space(text)
        style_id = _style_id(paragraph)
        if stripped == "参考文献":
            in_references = True
            continue
        if in_references and style_id == "Heading1":
            break
        if in_references:
            paragraphs.append((paragraph, text))
    return paragraphs


def _bibliography_diagnostics(
    xml_text: str,
    expected_ref_types_by_number: dict[int, str] | None = None,
) -> dict[str, object]:
    paragraphs = _bibliography_paragraphs_from_document_xml(xml_text)
    paragraph_texts = [text for _, text in paragraphs]
    ref_paragraphs = [
        (paragraph, ZOTERO_BIBLIOGRAPHY_ANCHOR_RE.sub("", text))
        for paragraph, text in paragraphs
        if re.match(r"^\[\d{1,3}\]", ZOTERO_BIBLIOGRAPHY_ANCHOR_RE.sub("", text))
    ]
    anchor_residues = [
        normalize_space(text)
        for text in paragraph_texts
        if ZOTERO_BIBLIOGRAPHY_ANCHOR_RE.search(text)
    ]
    stray_closing = [
        normalize_space(text)
        for text in paragraph_texts
        if normalize_space(text) == "]]"
    ]
    all_caps_author_samples: list[str] = []
    online_article_conference_samples: list[str] = []
    duplicate_doi_url_samples: list[str] = []
    access_field_samples: list[str] = []
    english_chinese_et_al_samples: list[str] = []
    report_standard_z_type_samples: list[str] = []
    report_standard_type_mismatch_samples: list[str] = []
    first_line_indent_count = 0
    hanging_indent_count = 0
    tight_line_spacing_count = 0
    non_left_alignment_count = 0
    expected_ref_types_by_number = expected_ref_types_by_number or {}

    for paragraph, text in ref_paragraphs:
        normalized = normalize_space(text)
        author_segment = normalized.split(".", 1)[0]
        if "," in author_segment and ALL_CAPS_ENGLISH_AUTHOR_RE.search(author_segment):
            all_caps_author_samples.append(normalized[:240])
        if re.match(r"^\[\d{1,3}\]\s*[A-Za-zÀ-ÿ]", normalized) and re.search(r"[,，]\s*等\b", author_segment):
            english_chinese_et_al_samples.append(normalized[:240])
        if re.search(r"\[[JC]/OL\]", normalized):
            online_article_conference_samples.append(normalized[:240])
        if re.search(r"\bDOI:\s*10\.", normalized, flags=re.IGNORECASE) or re.search(r"https?://", normalized):
            access_field_samples.append(normalized[:240])
        doi_match = re.search(r"DOI:\s*(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)", normalized, flags=re.IGNORECASE)
        resolver_match = DOI_RESOLVER_RE.search(normalized)
        if doi_match and resolver_match and doi_match.group(1).lower().rstrip(".") == resolver_match.group(1).lower().rstrip("."):
            duplicate_doi_url_samples.append(normalized[:240])

        number_match = re.match(r"^\[(\d{1,3})\]", normalized)
        expected_ref_type = expected_ref_types_by_number.get(int(number_match.group(1))) if number_match else None
        if expected_ref_type:
            expected_ref_type = expected_ref_type.upper()
            expected_code = "R" if expected_ref_type.startswith("R") else "S" if expected_ref_type.startswith("S") else None
            if expected_code:
                if re.search(r"\[Z(?:/OL)?\]", normalized):
                    report_standard_z_type_samples.append(normalized[:240])
                if not re.search(rf"\[{expected_code}(?:/OL)?\]", normalized):
                    report_standard_type_mismatch_samples.append(normalized[:240])

        ppr = _child(paragraph, "pPr")
        ind = _child(ppr, "ind") if ppr is not None else None
        if ind is not None:
            if (
                ind.attrib.get(f"{{{WORD_NS}}}firstLine")
                or ind.attrib.get(f"{{{WORD_NS}}}firstLineChars")
            ):
                first_line_indent_count += 1
            if ind.attrib.get(f"{{{WORD_NS}}}hanging"):
                hanging_indent_count += 1
        spacing = _child(ppr, "spacing") if ppr is not None else None
        if spacing is not None:
            try:
                line = int(spacing.attrib.get(f"{{{WORD_NS}}}line", "0"))
            except ValueError:
                line = 0
            line_rule = spacing.attrib.get(f"{{{WORD_NS}}}lineRule", "auto")
            if line_rule == "exact" and line and line < 360:
                tight_line_spacing_count += 1
        jc = _child(ppr, "jc") if ppr is not None else None
        if jc is not None and jc.attrib.get(f"{{{WORD_NS}}}val") not in {None, "left"}:
            non_left_alignment_count += 1

    return {
        "bibliography_reference_paragraphs": len(ref_paragraphs),
        "bibliography_anchor_residue_count": len(anchor_residues),
        "bibliography_anchor_residue_samples": anchor_residues[:10],
        "stray_closing_bracket_paragraphs": len(stray_closing),
        "stray_closing_bracket_samples": stray_closing[:10],
        "bibliography_all_caps_author_samples": all_caps_author_samples[:10],
        "bibliography_online_article_conference_samples": online_article_conference_samples[:10],
        "bibliography_duplicate_doi_url_samples": duplicate_doi_url_samples[:10],
        "bibliography_access_field_samples": access_field_samples[:10],
        "bibliography_english_chinese_et_al_count": len(english_chinese_et_al_samples),
        "bibliography_english_chinese_et_al_samples": english_chinese_et_al_samples[:10],
        "bibliography_report_standard_z_type_count": len(report_standard_z_type_samples),
        "bibliography_report_standard_z_type_samples": report_standard_z_type_samples[:10],
        "bibliography_report_standard_type_mismatch_count": len(report_standard_type_mismatch_samples),
        "bibliography_report_standard_type_mismatch_samples": report_standard_type_mismatch_samples[:10],
        "bibliography_first_line_indent_count": first_line_indent_count,
        "bibliography_hanging_indent_count": hanging_indent_count,
        "bibliography_tight_line_spacing_count": tight_line_spacing_count,
        "bibliography_non_left_alignment_count": non_left_alignment_count,
    }


def audit_zotero_docx(
    docx_path: Path,
    report_dir: Path = OUTPUT_DIR / "reports",
    checklist_path: Path | None = None,
) -> dict[str, object]:
    docx_path = Path(docx_path)
    report_dir.mkdir(parents=True, exist_ok=True)
    xml_files = list(_iter_docx_xml(docx_path))
    instructions = [
        instruction
        for _, xml in xml_files
        for instruction in _field_instructions(xml)
    ]
    field_structure: dict[str, int] = {
        "complex_field_begin_count": 0,
        "complex_field_separate_count": 0,
        "complex_field_end_count": 0,
        "complex_field_unclosed_count": 0,
        "complex_field_unmatched_end_count": 0,
        "bibliography_field_unclosed_count": 0,
        "visible_field_code_text_count": 0,
    }
    for _, xml in xml_files:
        diagnostics = _field_structure_diagnostics(xml)
        for key, value in diagnostics.items():
            field_structure[key] += value
    visible_paragraphs = [
        paragraph
        for name, xml in xml_files
        if _is_visible_story_xml(name)
        for paragraph in _visible_paragraphs_outside_fields(xml)
    ]
    expected_keys = _expected_zotero_key_sequence(checklist_path)
    reference_number_limit = _reference_number_limit_from_keys(expected_keys)
    static_markers = [
        marker
        for paragraph in visible_paragraphs
        for marker in CITATION_CLUSTER_RE.findall(paragraph)
        if _static_marker_within_reference_limit(marker, reference_number_limit)
    ]
    pandoc_markers, pandoc_keys = _pandoc_citation_items(visible_paragraphs)
    operator_placeholders = _operator_placeholders(visible_paragraphs)
    zotero_item_total, zotero_keys, zotero_parse_errors = _zotero_citation_items_from_instructions(instructions)
    expected_placeholders = _expected_operator_placeholders(checklist_path)
    document_xml = next((xml for name, xml in xml_files if name == "word/document.xml"), "")
    expected_ref_types_by_number = _expected_reference_types_by_number(checklist_path)
    bibliography_diagnostics = _bibliography_diagnostics(document_xml, expected_ref_types_by_number)
    audit = {
        "docx_path": display_path(docx_path),
        "xml_files_scanned": len(xml_files),
        "field_instructions": len(instructions),
        "zotero_item_fields": sum("ZOTERO_ITEM" in instruction for instruction in instructions),
        "csl_citation_fields": sum("CSL_CITATION" in instruction for instruction in instructions),
        "csl_bibliography_fields": sum("CSL_BIBLIOGRAPHY" in instruction for instruction in instructions),
        **field_structure,
        "extracted_citation_items": zotero_item_total,
        "zotero_citation_keys": zotero_keys,
        "zotero_citation_parse_errors": zotero_parse_errors,
        "zotero_citation_key_sequence_matches": bool(expected_keys) and zotero_keys == expected_keys,
        "pandoc_citation_markers": len(pandoc_markers),
        "pandoc_citation_items": len(pandoc_keys),
        "pandoc_citation_keys": pandoc_keys,
        "pandoc_citation_key_sequence_matches": bool(expected_keys) and pandoc_keys == expected_keys,
        "operator_placeholders": len(operator_placeholders),
        "operator_placeholder_samples": operator_placeholders[:20],
        "operator_placeholder_sequence_matches": bool(expected_placeholders) and operator_placeholders == expected_placeholders,
        "remaining_static_citation_markers": len(static_markers),
        "remaining_static_citation_samples": [normalize_space(item) for item in static_markers[:20]],
        **bibliography_diagnostics,
    }
    lines = [
        "# Zotero Word 字段审计",
        "",
        f"- DOCX: `{display_path(docx_path)}`",
        f"- XML files scanned: {audit['xml_files_scanned']}",
        f"- Field instructions parsed: {audit['field_instructions']}",
        f"- ZOTERO_ITEM fields: {audit['zotero_item_fields']}",
        f"- CSL_CITATION fields: {audit['csl_citation_fields']}",
        f"- CSL_BIBLIOGRAPHY fields: {audit['csl_bibliography_fields']}",
        f"- Complex field begins: {audit['complex_field_begin_count']}",
        f"- Complex field separates: {audit['complex_field_separate_count']}",
        f"- Complex field ends: {audit['complex_field_end_count']}",
        f"- Complex field unclosed: {audit['complex_field_unclosed_count']}",
        f"- Bibliography field unclosed: {audit['bibliography_field_unclosed_count']}",
        f"- Visible field-code text nodes: {audit['visible_field_code_text_count']}",
        f"- Extracted citation items: {audit['extracted_citation_items']}",
        f"- Zotero citation key sequence matches checklist: {audit['zotero_citation_key_sequence_matches']}",
        f"- Pandoc citation markers: {audit['pandoc_citation_markers']}",
        f"- Pandoc citation items: {audit['pandoc_citation_items']}",
        f"- Pandoc citation key sequence matches checklist: {audit['pandoc_citation_key_sequence_matches']}",
        f"- Operator placeholders: {audit['operator_placeholders']}",
        f"- Operator placeholder sequence matches checklist: {audit['operator_placeholder_sequence_matches']}",
        f"- Remaining static citation markers: {audit['remaining_static_citation_markers']}",
        f"- Bibliography anchor residues: {audit['bibliography_anchor_residue_count']}",
        f"- Stray closing bracket paragraphs: {audit['stray_closing_bracket_paragraphs']}",
        f"- Bibliography reference paragraphs: {audit['bibliography_reference_paragraphs']}",
        f"- Bibliography first-line indent paragraphs: {audit['bibliography_first_line_indent_count']}",
        f"- Bibliography hanging indent paragraphs: {audit['bibliography_hanging_indent_count']}",
        f"- Bibliography tight line-spacing paragraphs: {audit['bibliography_tight_line_spacing_count']}",
        f"- Bibliography non-left-aligned paragraphs: {audit['bibliography_non_left_alignment_count']}",
        f"- Bibliography all-caps author samples: {len(audit['bibliography_all_caps_author_samples'])}",
        f"- Bibliography [J/OL]/[C/OL] samples: {len(audit['bibliography_online_article_conference_samples'])}",
        f"- Bibliography duplicate DOI URL samples: {len(audit['bibliography_duplicate_doi_url_samples'])}",
        f"- Bibliography DOI/URL samples: {len(audit['bibliography_access_field_samples'])}",
        f"- Bibliography English entries using Chinese et-al term: {audit['bibliography_english_chinese_et_al_count']}",
        f"- Bibliography report/standard `[Z]` fallback samples: {audit['bibliography_report_standard_z_type_count']}",
        f"- Bibliography report/standard type mismatch samples: {audit['bibliography_report_standard_type_mismatch_count']}",
        "",
        "## Remaining Static Samples",
        "",
    ]
    if static_markers:
        lines.extend(f"- `{normalize_space(item)}`" for item in static_markers[:20])
    else:
        lines.append("- None detected.")
    lines.extend(["", "## Operator Placeholder Samples", ""])
    if operator_placeholders:
        lines.extend(f"- `{item}`" for item in operator_placeholders[:20])
    else:
        lines.append("- None detected.")
    for heading, key in [
        ("Bibliography Anchor Residues", "bibliography_anchor_residue_samples"),
        ("Stray Closing Bracket Samples", "stray_closing_bracket_samples"),
        ("All-Caps Author Samples", "bibliography_all_caps_author_samples"),
        ("[J/OL]/[C/OL] Samples", "bibliography_online_article_conference_samples"),
        ("Duplicate DOI URL Samples", "bibliography_duplicate_doi_url_samples"),
        ("DOI/URL Samples", "bibliography_access_field_samples"),
        ("English Entries Using Chinese Et-Al Term", "bibliography_english_chinese_et_al_samples"),
        ("Report/Standard `[Z]` Fallback Samples", "bibliography_report_standard_z_type_samples"),
        ("Report/Standard Type Mismatch Samples", "bibliography_report_standard_type_mismatch_samples"),
    ]:
        lines.extend(["", f"## {heading}", ""])
        samples = audit.get(key, [])
        if samples:
            lines.extend(f"- `{item}`" for item in samples)
        else:
            lines.append("- None detected.")
    lines.extend(["", "## Parsed Citation Keys", ""])
    if zotero_keys:
        lines.extend(f"- `{item}`" for item in zotero_keys[:50])
    else:
        lines.append("- None detected.")
    (report_dir / "zotero_field_audit.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return audit


def _set_attr(element: ElementTree.Element, name: str, value: str) -> None:
    element.set(f"{{{WORD_NS}}}{name}", value)


def _remove_children(parent: ElementTree.Element, tag: str) -> None:
    for child in list(parent):
        if child.tag == f"{{{WORD_NS}}}{tag}":
            parent.remove(child)


PARAGRAPH_PROPERTY_ORDER = {
    "pStyle": 10,
    "keepNext": 20,
    "keepLines": 30,
    "pageBreakBefore": 40,
    "framePr": 50,
    "widowControl": 60,
    "numPr": 70,
    "suppressLineNumbers": 80,
    "pBdr": 90,
    "shd": 100,
    "tabs": 110,
    "suppressAutoHyphens": 120,
    "kinsoku": 130,
    "wordWrap": 140,
    "overflowPunct": 150,
    "topLinePunct": 160,
    "autoSpaceDE": 170,
    "autoSpaceDN": 180,
    "bidi": 190,
    "adjustRightInd": 200,
    "snapToGrid": 210,
    "spacing": 220,
    "ind": 230,
    "contextualSpacing": 240,
    "mirrorIndents": 250,
    "suppressOverlap": 260,
    "jc": 270,
    "textDirection": 280,
    "textAlignment": 290,
    "textboxTightWrap": 300,
    "outlineLvl": 310,
    "divId": 320,
    "cnfStyle": 330,
    "rPr": 340,
    "sectPr": 350,
    "pPrChange": 360,
}

WORD_IGNORABLE_NAMESPACE_URIS = {
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    "w15": "http://schemas.microsoft.com/office/word/2012/wordml",
    "w16se": "http://schemas.microsoft.com/office/word/2015/wordml/symex",
    "w16cid": "http://schemas.microsoft.com/office/word/2016/wordml/cid",
    "w16": "http://schemas.microsoft.com/office/word/2018/wordml",
    "w16cex": "http://schemas.microsoft.com/office/word/2018/wordml/cex",
    "w16sdtdh": "http://schemas.microsoft.com/office/word/2020/wordml/sdtdatahash",
    "w16sdtfl": "http://schemas.microsoft.com/office/word/2024/wordml/sdtformatlock",
    "w16du": "http://schemas.microsoft.com/office/word/2023/wordml/word16du",
    "wp14": "http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing",
}


def _order_paragraph_properties(ppr: ElementTree.Element) -> None:
    indexed_children = list(enumerate(list(ppr)))
    indexed_children.sort(
        key=lambda item: (
            PARAGRAPH_PROPERTY_ORDER.get(item[1].tag.rsplit("}", 1)[-1], 1000),
            item[0],
        )
    )
    ppr[:] = [child for _, child in indexed_children]


def _xml_namespace_map(xml_text: str) -> dict[str, str]:
    namespace_map: dict[str, str] = {}
    for _, namespace in ElementTree.iterparse(io.StringIO(xml_text), events=("start-ns",)):
        prefix, uri = namespace
        namespace_map[prefix] = uri
    return namespace_map


def _register_xml_namespaces(xml_text: str) -> dict[str, str]:
    namespace_map = _xml_namespace_map(xml_text)
    for prefix, uri in namespace_map.items():
        if prefix != "xml" and not re.fullmatch(r"ns\d+", prefix):
            ElementTree.register_namespace(prefix, uri)
    return namespace_map


def normalize_latex_artifacts(text: str) -> str:
    return LATEX_SUPERSCRIPT_RE.sub(
        lambda match: match.group(1) + match.group(2).translate(SUPERSCRIPT_DIGITS),
        text,
    )


def normalize_bibliography_entry_text(text: str) -> str:
    text = normalize_latex_artifacts(text).replace("\t", " ")
    text = URL_TEXT_RE.sub("", text)
    text = DOI_TEXT_RE.sub("", text)
    text = re.sub(r"^(\[\d{1,3}\])\s+", r"\1 ", text)
    text = re.sub(r"\s+([,.;:])", r"\1", text)
    text = re.sub(r"\s{2,}", " ", text).strip()
    if text and not re.search(r"[.。]$", text):
        text += "."
    return text


def _is_field_run(run: ElementTree.Element) -> bool:
    return (
        run.find("w:fldChar", {"w": WORD_NS}) is not None
        or run.find("w:instrText", {"w": WORD_NS}) is not None
    )


def _field_run_type(run: ElementTree.Element) -> str | None:
    fld_char = run.find("w:fldChar", {"w": WORD_NS})
    if fld_char is None:
        return None
    return fld_char.attrib.get(f"{{{WORD_NS}}}fldCharType")


def _normalize_bibliography_entry_runs(paragraph: ElementTree.Element) -> None:
    cleaned_text = normalize_bibliography_entry_text(_paragraph_text_with_tabs(paragraph))
    if not cleaned_text:
        return

    for child in list(paragraph):
        if child.tag == f"{{{WORD_NS}}}r" and not _is_field_run(child):
            paragraph.remove(child)

    insert_at = len(list(paragraph))
    for index, child in enumerate(list(paragraph)):
        if child.tag == f"{{{WORD_NS}}}r" and _field_run_type(child) == "end":
            insert_at = index
            break

    run = ElementTree.Element(f"{{{WORD_NS}}}r")
    text_element = ElementTree.SubElement(run, f"{{{WORD_NS}}}t")
    text_element.text = cleaned_text
    paragraph.insert(insert_at, run)


def _ensure_ignorable_namespace_declarations(xml_text: str, namespace_map: dict[str, str]) -> str:
    root_match = re.search(r"<[^!?][^>]*>", xml_text)
    if not root_match:
        return xml_text
    root_start = root_match.group(0)
    ignorable_match = re.search(r":Ignorable=\"([^\"]+)\"", root_start)
    if not ignorable_match:
        return xml_text
    additions = []
    for prefix in ignorable_match.group(1).split():
        uri = namespace_map.get(prefix) or WORD_IGNORABLE_NAMESPACE_URIS.get(prefix)
        if uri and f"xmlns:{prefix}=" not in root_start:
            additions.append(f' xmlns:{prefix}="{uri}"')
    if not additions:
        return xml_text
    replacement = root_start[:-1] + "".join(additions) + ">"
    return xml_text[: root_match.start()] + replacement + xml_text[root_match.end() :]


def _apply_bibliography_paragraph_format(paragraph: ElementTree.Element) -> None:
    ppr = _ensure_child(paragraph, "pPr")
    pstyle = _ensure_child(ppr, "pStyle")
    _set_attr(pstyle, "val", "Bibliography")

    _remove_children(ppr, "tabs")
    tabs = ElementTree.SubElement(ppr, f"{{{WORD_NS}}}tabs")
    tab = ElementTree.SubElement(tabs, f"{{{WORD_NS}}}tab")
    _set_attr(tab, "val", "left")
    _set_attr(tab, "pos", BIBLIOGRAPHY_HANGING_TWIPS)

    spacing = _ensure_child(ppr, "spacing")
    _set_attr(spacing, "before", "0")
    _set_attr(spacing, "after", "0")
    _set_attr(spacing, "line", BIBLIOGRAPHY_LINE_SPACING)
    _set_attr(spacing, "lineRule", "auto")

    jc = _ensure_child(ppr, "jc")
    _set_attr(jc, "val", "left")

    ind = _ensure_child(ppr, "ind")
    for key in ("firstLine", "firstLineChars"):
        ind.attrib.pop(f"{{{WORD_NS}}}{key}", None)
    _set_attr(ind, "left", BIBLIOGRAPHY_HANGING_TWIPS)
    _set_attr(ind, "hanging", BIBLIOGRAPHY_HANGING_TWIPS)

    rpr = _ensure_child(ppr, "rPr")
    rfonts = _ensure_child(rpr, "rFonts")
    _set_attr(rfonts, "ascii", "Times New Roman")
    _set_attr(rfonts, "hAnsi", "Times New Roman")
    _set_attr(rfonts, "eastAsia", "SimSun")
    _set_attr(rfonts, "cs", "Times New Roman")
    sz = _ensure_child(rpr, "sz")
    _set_attr(sz, "val", "24")
    szcs = _ensure_child(rpr, "szCs")
    _set_attr(szcs, "val", "24")

    for run in paragraph.findall("w:r", {"w": WORD_NS}):
        run_rpr = _ensure_run_properties(run)
        run_fonts = _ensure_child(run_rpr, "rFonts")
        _set_attr(run_fonts, "ascii", "Times New Roman")
        _set_attr(run_fonts, "hAnsi", "Times New Roman")
        _set_attr(run_fonts, "eastAsia", "SimSun")
        _set_attr(run_fonts, "cs", "Times New Roman")
        run_sz = _ensure_child(run_rpr, "sz")
        _set_attr(run_sz, "val", "24")
        run_szcs = _ensure_child(run_rpr, "szCs")
        _set_attr(run_szcs, "val", "24")
    _order_paragraph_properties(ppr)


def _complex_field_end_run() -> ElementTree.Element:
    run = ElementTree.Element(f"{{{WORD_NS}}}r")
    field_end = ElementTree.SubElement(run, f"{{{WORD_NS}}}fldChar")
    _set_attr(field_end, "fldCharType", "end")
    return run


def _postprocess_document_xml_for_bibliography(xml_text: str) -> str:
    namespace_map = _register_xml_namespaces(xml_text)
    xml_text = normalize_latex_artifacts(xml_text)
    xml_text = ZOTERO_BIBLIOGRAPHY_ANCHOR_RE.sub("", xml_text)
    root = ElementTree.fromstring(xml_text.encode("utf-8"))
    parent_by_child = {child: parent for parent in root.iter() for child in parent}
    in_references = False
    last_bibliography_paragraph: ElementTree.Element | None = None
    for paragraph in list(root.iter(f"{{{WORD_NS}}}p")):
        text = _paragraph_text_with_tabs(paragraph)
        stripped = normalize_space(text)
        style_id = _style_id(paragraph)
        if stripped == "参考文献":
            in_references = True
            continue
        if in_references and style_id == "Heading1":
            break
        if not in_references:
            continue
        if stripped == "]]":
            if paragraph.findall(".//w:fldChar", {"w": WORD_NS}):
                for text_element in paragraph.findall(".//w:t", {"w": WORD_NS}):
                    text_element.text = (text_element.text or "").replace("]]", "")
                continue
            parent = parent_by_child.get(paragraph)
            if parent is not None:
                parent.remove(paragraph)
            continue
        if re.match(r"^\[\d{1,3}\]", text):
            last_bibliography_paragraph = paragraph
            _normalize_bibliography_entry_runs(paragraph)
            _apply_bibliography_paragraph_format(paragraph)
    diagnostics = _field_structure_diagnostics(ElementTree.tostring(root, encoding="unicode"))
    if (
        last_bibliography_paragraph is not None
        and diagnostics["complex_field_unclosed_count"] == 1
        and diagnostics["bibliography_field_unclosed_count"] == 1
        and diagnostics["complex_field_unmatched_end_count"] == 0
    ):
        parent = parent_by_child.get(last_bibliography_paragraph)
        if parent is not None:
            field_end_paragraph = ElementTree.Element(f"{{{WORD_NS}}}p")
            field_end_paragraph.append(_complex_field_end_run())
            parent.insert(list(parent).index(last_bibliography_paragraph) + 1, field_end_paragraph)
    output = ElementTree.tostring(root, encoding="unicode", xml_declaration=True)
    return _ensure_ignorable_namespace_declarations(output, namespace_map)


def _postprocess_styles_xml_for_bibliography(xml_text: str) -> str:
    namespace_map = _register_xml_namespaces(xml_text)
    root = ElementTree.fromstring(xml_text.encode("utf-8"))
    style = root.find(".//w:style[@w:styleId='Bibliography']", {"w": WORD_NS})
    if style is None:
        style = ElementTree.SubElement(root, f"{{{WORD_NS}}}style")
        _set_attr(style, "type", "paragraph")
        _set_attr(style, "styleId", "Bibliography")
        name = ElementTree.SubElement(style, f"{{{WORD_NS}}}name")
        _set_attr(name, "val", "Bibliography")
    ppr = _ensure_child(style, "pPr")
    probe = ElementTree.Element(f"{{{WORD_NS}}}p")
    probe.append(ppr)
    _apply_bibliography_paragraph_format(probe)
    if probe.find(f"{{{WORD_NS}}}pPr") is None:
        style.append(ppr)
    output = ElementTree.tostring(root, encoding="unicode", xml_declaration=True)
    return _ensure_ignorable_namespace_declarations(output, namespace_map)


def postprocess_zotero_bibliography_docx(input_docx: Path, output_docx: Path | None = None) -> Path:
    """Clean Zotero bibliography anchor residue and normalize reference paragraph format."""
    input_docx = Path(input_docx)
    output_docx = Path(output_docx) if output_docx else input_docx
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    temp_output = output_docx.with_name(f".{output_docx.stem}.{uuid.uuid4().hex}.docx")
    with zipfile.ZipFile(input_docx, "r") as source, zipfile.ZipFile(temp_output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                data = _postprocess_document_xml_for_bibliography(data.decode("utf-8")).encode("utf-8")
            elif info.filename == "word/styles.xml":
                data = _postprocess_styles_xml_for_bibliography(data.decode("utf-8")).encode("utf-8")
            target.writestr(info, data)
    shutil.move(str(temp_output), output_docx)
    return output_docx


def optimize_docx_media(input_docx: Path, media_source_docx: Path, output_docx: Path | None = None) -> Path:
    """Replace Word media with matching files from a source DOCX and deflate the package."""
    input_docx = Path(input_docx)
    media_source_docx = Path(media_source_docx)
    output_docx = Path(output_docx) if output_docx else input_docx
    output_docx.parent.mkdir(parents=True, exist_ok=True)

    source_media: dict[str, bytes] = {}
    with zipfile.ZipFile(media_source_docx, "r") as media_source:
        for info in media_source.infolist():
            if info.filename.startswith("word/media/"):
                source_media[info.filename] = media_source.read(info.filename)

    temp_output = output_docx.with_name(f".{output_docx.stem}.{uuid.uuid4().hex}.docx")
    with zipfile.ZipFile(input_docx, "r") as source, zipfile.ZipFile(
        temp_output,
        "w",
        zipfile.ZIP_DEFLATED,
        compresslevel=6,
    ) as target:
        for info in source.infolist():
            data = source_media.get(info.filename)
            if data is None:
                data = source.read(info.filename)
            target.writestr(info.filename, data)
    shutil.move(str(temp_output), output_docx)
    return output_docx


def replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.text == text:
        return
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def insert_paragraph_after(paragraph, text: str, style: str | None = None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    new_para.add_run(text)
    return new_para


def formatted_reference(ref: Reference, global_id: int | None = None) -> str:
    gid = global_id if global_id is not None else ref.global_id
    body = normalize_space(ref.raw_text)
    body = re.sub(r"\bdoi:", "DOI:", body, flags=re.IGNORECASE)
    if not body.endswith("."):
        body += "."
    return f"[{gid}] {body}"


def citation_replacer_for_model(model: ReferenceModel):
    refs_by_context = {
        (ref.source_chapter, ref.source_section, ref.local_id): ref
        for ref in model.references
    }

    def replace_marker(source_chapter: str, source_section: str, text: str) -> str:
        def repl(match: re.Match[str]) -> str:
            local_id = int(match.group(1))
            ref = refs_by_context.get((source_chapter, source_section, local_id))
            if not ref:
                return match.group(0)
            return f"[{model.global_id_by_ref_id[ref.ref_id]}]"

        return compress_citation_runs(CITATION_RE.sub(repl, text))

    return replace_marker


def citation_replacer_by_paragraph(model: ReferenceModel):
    occurrences_by_paragraph: dict[int, dict[int, CitationOccurrence]] = {}
    for occurrence in model.citation_occurrences:
        occurrences_by_paragraph.setdefault(occurrence.paragraph_index, {})[occurrence.marker_index] = occurrence

    def replace_text(paragraph_index: int, text: str) -> str:
        occurrence_by_marker = occurrences_by_paragraph.get(paragraph_index, {})
        if not occurrence_by_marker:
            return text

        marker_index = -1

        def repl(match: re.Match[str]) -> str:
            nonlocal marker_index
            marker_index += 1
            occurrence = occurrence_by_marker.get(marker_index)
            if occurrence is None:
                return match.group(0)
            global_id = model.global_id_by_ref_id.get(occurrence.ref_id)
            if global_id is None:
                return match.group(0)
            return f"[{global_id}]"

        return compress_citation_runs(CITATION_RE.sub(repl, text))

    return replace_text


def rewrite_main_docx(
    model: ReferenceModel,
    output_path: Path | None = None,
    input_docx_path: Path | str | None = None,
) -> Path:
    output_path = Path(output_path or UNIFIED_DOC)
    input_path = Path(input_docx_path or model.input_docx_path or INPUT_DOC)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    working_path = output_path.parent / f".{output_path.stem}.{uuid.uuid4().hex}.tmp.docx"
    shutil.copyfile(input_path, working_path)
    doc = Document(working_path)
    replace_text = citation_replacer_by_paragraph(model)

    for paragraph_index in sorted({occ.paragraph_index for occ in model.citation_occurrences}):
        if paragraph_index >= len(doc.paragraphs):
            continue
        paragraph = doc.paragraphs[paragraph_index]
        if is_heading(paragraph) or "[" not in paragraph.text:
            continue
        replace_paragraph_text(paragraph, replace_text(paragraph_index, paragraph.text))

    for idx in sorted(model.reference_paragraphs or model.chapter_one_reference_paragraphs, reverse=True):
        if idx >= len(doc.paragraphs):
            continue
        delete_paragraph(doc.paragraphs[idx])

    ref_heading_idx = final_bibliography_heading_index(doc)
    if ref_heading_idx is None:
        doc.add_heading("参考文献", level=1)
        ref_heading_idx = len(doc.paragraphs) - 1
    insert_idx = ref_heading_idx + 1
    while insert_idx < len(doc.paragraphs):
        paragraph = doc.paragraphs[insert_idx]
        if paragraph.style.name == "Heading 1":
            break
        delete_paragraph(paragraph)

    anchor = doc.paragraphs[ref_heading_idx]
    for ref in model.unique_references:
        anchor = insert_paragraph_after(anchor, formatted_reference(ref, ref.global_id), style="Normal")

    doc.save(working_path)
    shutil.move(str(working_path), output_path)
    build_copy = BUILD_DIR / "thesis_v2.docx"
    if output_path.resolve() != build_copy.resolve():
        shutil.copyfile(output_path, build_copy)
    return output_path


def ensure_build_dirs() -> None:
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "doc").mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "zotero").mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "reports").mkdir(parents=True, exist_ok=True)


def write_json(path: Path, data: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def write_raw_refs(model: ReferenceModel) -> None:
    ensure_build_dirs()
    write_json(BUILD_DIR / "raw_refs.json", [asdict(ref) for ref in model.references])
    write_json(BUILD_DIR / "citation_occurrences.json", [asdict(occ) for occ in model.citation_occurrences])


def write_ref_map(model: ReferenceModel) -> None:
    ensure_build_dirs()
    data = {
        "global_id_by_ref_id": model.global_id_by_ref_id,
        "unique_references": [asdict(ref) for ref in model.unique_references],
    }
    write_json(BUILD_DIR / "ref_map.json", data)
    lines = ["# Reference Map", ""]
    lines.append("| Global ID | Source | Local ID | Title | Raw Text |")
    lines.append("|---:|---|---:|---|---|")
    for ref in sorted(model.references, key=lambda item: (item.global_id or 9999, item.source_chapter, item.local_id)):
        raw = ref.raw_text.replace("|", "\\|")
        title = ref.title.replace("|", "\\|")
        lines.append(f"| {ref.global_id} | {ref.source_section} | {ref.local_id} | {title} | {raw} |")
    (OUTPUT_DIR / "reports").mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "reports" / "reference_map.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "reference_map"
    sheet.append(["global_id", "source_chapter", "source_section", "local_id", "title", "raw_text", "doi", "url"])
    for ref in sorted(model.references, key=lambda item: (item.global_id or 9999, item.source_chapter, item.local_id)):
        sheet.append([ref.global_id, ref.source_chapter, ref.source_section, ref.local_id, ref.title, ref.raw_text, ref.doi, ref.url])
    workbook.save(OUTPUT_DIR / "reports" / "reference_map.xlsx")


def write_final_references(model: ReferenceModel) -> None:
    ensure_build_dirs()
    lines = [formatted_reference(ref, ref.global_id) for ref in model.unique_references]
    (BUILD_DIR / "refs_final.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")


def bibtex_key(ref: Reference) -> str:
    return f"ref{int(ref.global_id or 0):03d}"


def bibtex_escape(text: str | None) -> str:
    if not text:
        return ""
    return text.replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}")


def is_corporate_author(part: str) -> bool:
    return zotero_is_corporate_author(part)


def split_author_parts(value: str) -> tuple[list[str], bool]:
    return zotero_split_author_parts(value)


def bibtex_author_value(value: str | None) -> str:
    if not value:
        return ""
    return creators_to_bibtex(creators_from_responsibility(value))


def bibtex_type(ref: Reference) -> str:
    if ref.ref_type == "J":
        return "article"
    if ref.ref_type == "C":
        return "inproceedings"
    return "misc"


def reference_container_fields(ref: Reference) -> dict[str, str]:
    text = normalize_space(ref.raw_text)
    fields: dict[str, str] = {}

    journal_match = re.search(
        r"\[J(?:/OL)?\]\.\s*(?P<journal>.+?),\s*(?P<year>(?:19|20)\d{2})"
        r"(?:,\s*(?P<volume>[^():，,]+)(?:\((?P<issue>[^)]+)\))?)?"
        r"(?::\s*(?P<pages>[^.。]+))?",
        text,
    )
    if journal_match:
        fields["journal"] = normalize_space(journal_match.group("journal"))
        if journal_match.group("volume"):
            fields["volume"] = normalize_space(journal_match.group("volume"))
        if journal_match.group("issue"):
            fields["number"] = normalize_space(journal_match.group("issue"))
        if journal_match.group("pages"):
            fields["pages"] = normalize_space(journal_match.group("pages"))
        return fields

    conference_match = re.search(
        r"\[C(?:/OL)?\]//(?P<booktitle>.+?)\.\s*"
        r"(?:(?P<publisher>[^,，.。]+),\s*)?(?P<year>(?:19|20)\d{2})"
        r"(?::\s*(?P<pages>[^.。]+))?",
        text,
    )
    if conference_match:
        fields["booktitle"] = normalize_space(conference_match.group("booktitle"))
        if conference_match.group("publisher"):
            fields["publisher"] = normalize_space(conference_match.group("publisher"))
        if conference_match.group("pages"):
            fields["pages"] = normalize_space(conference_match.group("pages"))
    return fields


def write_bibtex(model: ReferenceModel) -> None:
    ensure_build_dirs()
    entries: list[str] = []
    for ref in model.unique_references:
        fields = {
            "author": bibtex_author_value(ref.responsibility),
            "title": ref.title or ref.raw_text,
            "year": ref.year,
            **reference_container_fields(ref),
            "doi": ref.doi,
            "url": ref.url,
            "note": f"Global reference [{ref.global_id}]; source {ref.source_section} local [{ref.local_id}]",
        }
        body_lines = []
        for key, value in fields.items():
            if value:
                escaped = value.replace("\\", "\\\\") if key == "author" else bibtex_escape(value)
                body_lines.append(f"  {key} = {{{escaped}}},")
        entries.append(f"@{bibtex_type(ref)}{{{bibtex_key(ref)},\n" + "\n".join(body_lines) + "\n}")
    (OUTPUT_DIR / "zotero" / "refs.bib").write_text("\n\n".join(entries) + "\n", encoding="utf-8")


def csl_type(ref: Reference) -> str:
    if ref.ref_type == "J":
        return "article-journal"
    if ref.ref_type == "C":
        return "paper-conference"
    if ref.ref_type in {"R/OL", "R"}:
        return "report"
    if ref.ref_type in {"S/OL", "S"}:
        return "standard"
    return "webpage"


def literal_author(value: str) -> list[dict[str, str]]:
    return creators_to_csl(creators_from_responsibility(value))


def write_csl_json(model: ReferenceModel) -> None:
    ensure_build_dirs()
    items = []
    for ref in model.unique_references:
        container_fields = reference_container_fields(ref)
        item = {
            "id": bibtex_key(ref),
            "type": csl_type(ref),
            "title": ref.title or ref.raw_text,
            "author": literal_author(ref.responsibility),
            "note": f"Global reference [{ref.global_id}]; source {ref.source_section} local [{ref.local_id}]",
        }
        if ref.year:
            item["issued"] = {"date-parts": [[int(ref.year)]]}
        if ref.doi:
            item["DOI"] = ref.doi
        if ref.url:
            item["URL"] = ref.url
        if container_fields.get("journal"):
            item["container-title"] = container_fields["journal"]
        if container_fields.get("booktitle"):
            item["container-title"] = container_fields["booktitle"]
        if container_fields.get("volume"):
            item["volume"] = container_fields["volume"]
        if container_fields.get("number"):
            item["issue"] = container_fields["number"]
        if container_fields.get("pages"):
            item["page"] = container_fields["pages"]
        if container_fields.get("publisher"):
            item["publisher"] = container_fields["publisher"]
        items.append(item)
    write_json(OUTPUT_DIR / "zotero" / "refs.csl.json", items)


def write_zotero_group_handoff(
    model: ReferenceModel,
    output_dir: Path | None = None,
) -> Path:
    output_dir = Path(output_dir or OUTPUT_DIR / "zotero")
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "group_library_handoff.md"
    content = f"""# Zotero Group Library 交接清单

本清单用于把静态统一编号版本迁移到 Zotero Word 动态引用。推荐使用 Zotero Group Library，而不是个人 My Library，这样朋友打开同一个 Word 后可以从共享组库继续 Add/Edit Citation、Refresh 和更新 bibliography。

## 输入文件

- `output/zotero/refs.bib`: 全文 {len(model.unique_references)} 条唯一文献，可导入 Zotero Group Library。
- `output/zotero/refs.csl.json`: 备用导入格式。
- `output/zotero/migration_checklist.md`: 正文静态引用到 Zotero 条目的逐位置迁移清单。
- `output/reports/reference_map.xlsx`: 统一编号、原章节编号与题名映射。

## 准备步骤

1. 在 Zotero 中创建或打开共享 Group Library，并邀请朋友加入。
2. 在该 Group Library 下导入 `refs.bib`，确认条目数为 {len(model.unique_references)}。
3. 安装并选择 `Chinese Std GB/T 7714-2015 (numeric)` 或学校认可的 GB/T 7714-2015 顺序编码制 CSL。
4. 打开 `output/doc/论文_统一编号.docx` 的副本，使用 Zotero Word 插件设置 Document Preferences 为 GB/T 7714-2015 numeric。
5. 按 `migration_checklist.md` 逐个替换正文静态 `[N]`；同一位置有多个 key 时，在同一个 Zotero citation 中加入多篇文献。
6. 每处理 5-10 个位置保存一个 checkpoint 副本。
7. 全部正文引用替换后，删除静态文末参考文献，用 Zotero Add/Edit Bibliography 生成动态 bibliography，再运行 Refresh。

## 注意

- 不要直接编辑 `.docx` XML 来伪造 Zotero 字段；字段必须由 Zotero Word 插件创建。
- 如果从个人 My Library 插入引用，朋友电脑上可能变成 orphaned items，不利于后续协作维护。
- Computer Use 只能代替点击和输入，遇到 Zotero 登录、组库权限、Word 宏权限弹窗时需要人工确认。
"""
    path.write_text(content, encoding="utf-8")
    return path


def verification_rows(model: ReferenceModel) -> list[dict[str, str | int | None]]:
    rows = []
    for ref in model.references:
        rows.append(
            {
                "global_id": ref.global_id,
                "source_section": ref.source_section,
                "local_id": ref.local_id,
                "title": ref.title,
                "status": ref.verification_status,
                "evidence": ref.verification_evidence or "No public evidence attached by automated verification.",
                "note": ref.verification_note,
                "raw_text": ref.raw_text,
            }
        )
    return rows


def write_verification_report(model: ReferenceModel) -> None:
    ensure_build_dirs()
    rows = verification_rows(model)
    lines = [
        "# 参考文献核查报告",
        "",
        "状态说明：`manual_review_needed` 表示源文档没有 DOI/URL，脚本无法只靠本地字段确认；`needs_check` 表示脚本给出候选官方核验入口。",
        "",
        "| Global ID | Section | Local ID | Status | Evidence | Note | Title |",
        "|---:|---|---:|---|---|---|---|",
    ]
    for row in rows:
        lines.append(
            f"| {row['global_id']} | {row['source_section']} | {row['local_id']} | "
            f"{row['status']} | {str(row['evidence']).replace('|', '/')} | "
            f"{str(row['note']).replace('|', '/')} | {str(row['title']).replace('|', '/')} |"
        )
    lines.append("")
    lines.append("## 原始条目")
    lines.append("")
    for row in rows:
        lines.append(f"- [{row['global_id']}] {row['raw_text']} ({row['status']})")
    if model.missing_reference_blocks:
        lines.extend(["", "## Missing Reference Blocks", ""])
        for item in model.missing_reference_blocks:
            markers = ", ".join(str(marker) for marker in item.get("citation_markers", []))
            lines.append(f"- `{item.get('source_chapter')}`: {item.get('issue')} ({markers})")
    content = "\n".join(lines) + "\n"
    (BUILD_DIR / "verify_report.md").write_text(content, encoding="utf-8")
    (OUTPUT_DIR / "reports" / "verify_report.md").write_text(content, encoding="utf-8")


def write_zotero_readme(model: ReferenceModel) -> None:
    ensure_build_dirs()
    content = f"""# Zotero 接手说明

本阶段交付的是静态统一编号 Word。正文中的 `[N]` 是普通文本，不是 Zotero 动态域。

## 文件

- `output/doc/论文_统一编号.docx`: 可直接提交和人工编辑的 Word。
- `output/zotero/refs.bib`: Zotero 可导入的全文参考文献库，共 {len(model.unique_references)} 条唯一文献。
- `output/zotero/refs.csl.json`: CSL JSON 备用导入格式。
- `output/reports/reference_map.xlsx`: 统一编号、原章节编号与题名映射。
- `output/reports/verify_report.md`: 参考文献核查报告。
- `output/zotero/migration_checklist.md`: 静态 `[N]` 到 Zotero 条目的逐位置迁移清单。
- `output/doc/论文_zotero_dynamic_operator.docx`: 官方 Word 插件 UI 兜底操作版，正文引用为唯一 `[[ZOTERO_Pxxx]]` 占位符。
- `output/zotero/group_library_handoff.md`: Zotero Group Library 交接清单。

## 朋友电脑操作（推荐 Group Library）

1. 安装 Zotero 和 Zotero Word 插件。
2. 创建或加入共享 Group Library，把 `refs.bib` 导入该组库。
3. 安装或选择 Chinese Std GB/T 7714-2015 numeric 样式（GB/T 7714-2015 顺序编码制）。
4. 后续新增引用时，从同一个 Group Library 用 Zotero Word 插件 Add/Edit Citation 插入。
5. 如果从个人 My Library 插入，朋友端可能只能看到 orphaned items，不利于协作刷新。

## 切换到 Zotero 全托管

1. 打开 `output/zotero/migration_checklist.md`，按正文段落顺序处理每个 `Position`。
2. 在 Word 中删除该位置静态 `[N]`，用 Zotero 插件插入 `Zotero keys` 对应条目；连续引用如 `[27-29]` 在同一个 citation 里加入多篇文献。
3. 每处理 5-10 个位置保存一个 checkpoint 副本，便于 Word/Zotero 弹窗或误操作后回退。
4. 全部替换后删除静态文末参考文献，使用 Zotero Add/Edit Bibliography 生成。
5. 之后新增、删除、移动引用时，用 Zotero Refresh 统一刷新编号和文末清单。

## 自动化边界

CLI 工具负责生成迁移清单、维护 `refs.bib`/`refs.csl.json`、准备 operator 输入副本，并审计 docx 是否包含真实 Zotero 字段。最终交付优先使用官方 Word + Zotero 插件流程；直接生成 Zotero Word 字段只允许在 `build/`、`output/`、`tmp/` 实验副本中进行，且必须通过 Word 打开无修复、Zotero Refresh、重建 bibliography 和映射审计后才可交付。
"""
    (OUTPUT_DIR / "zotero" / "README_zotero.md").write_text(content, encoding="utf-8")


def run_all_exports(input_docx: Path | str | None = None, output_root: Path | str | None = None) -> ReferenceModel:
    input_path = configure_paths(input_docx, output_root)
    model = build_reference_model(input_path)
    write_raw_refs(model)
    write_ref_map(model)
    write_verification_report(model)
    write_final_references(model)
    write_bibtex(model)
    write_csl_json(model)
    rewrite_main_docx(model)
    write_zotero_migration_checklist(model)
    write_zotero_group_handoff(model)
    write_zotero_readme(model)
    return model
